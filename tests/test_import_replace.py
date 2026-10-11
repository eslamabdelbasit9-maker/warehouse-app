"""استيراد ملف موقع مرتين: التانية بـ «استبدال» تمسح القديم وتحط الجديد وترجّع ربط الفواتير المرفوعة."""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

import tests.test_flow  # noqa: F401  (يضبط قاعدة بيانات الاختبار)
from app.db import SessionLocal
from app.models import OpeningBalance, Receipt, Request, Site
from app.main import init_db
from scripts.import_excel import run


def book(extra_issue=False, receipts=1):
    wb = Workbook()
    b = wb.active
    b.title = "الرصيد"
    b.append(["رصيد"])
    b.append(["كود الصنف", "اسم الصنف", "الموقع", "الرصيد الافتتاحي", "قيمة الرصيد الافتتاحي", "وحدة القياس"])
    b.append(["SP_T1", "صنف تجربة", "تست", 5, 50, "قطعة"])
    b.add_table(Table(displayName="BalanceTable", ref="A2:F3"))
    r = wb.create_sheet("الوارد")
    r.append(["التاريخ", "رقم الفاتورة", "كود الصنف", "اسم الصنف", "الكمية", "سعر الوحدة", "الإجمالي", "المورد", "مرفق الفاتورة"])
    for i in range(receipts):
        r.append([date(2026, 9, 1 + i), f"INV-{i + 1}", "SP_T1", "صنف تجربة", 2, 10, 20, "مورد", None])
    s = wb.create_sheet("الصرف")
    s.append(["رقم الطلب", "التاريخ", "كود الصنف", "اسم الصنف", "الكمية", "المشروع", "طالب الصرف", "السبب",
              "حالة الاعتماد", "المهندس المعتمد", "سعر الوحدة", "القيمة", "الشهر", "ReqID"])
    rows = [["REQ-2609010801-SP_T1", "2026-09-01", "SP_T1", "صنف تجربة", 1, "كسارة", "x", "y", "معتمد نهائي", None, 10, 10, "2026-09", "REQ-2609010801"]]
    if extra_issue:
        rows.append(["REQ-2610010801-SP_T1", "2026-10-01", "SP_T1", "صنف تجربة", 1, "كسارة", "x", "y", "معتمد نهائي", None, 10, 10, "2026-10", "REQ-2610010801"])
    for x in rows:
        s.append(x)
    s.add_table(Table(displayName="IssueTable", ref=f"A1:N{len(rows) + 1}"))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_import_replace():
    init_db()
    run(book(), "TST", "تست")
    db = SessionLocal()
    site = db.query(Site).filter_by(code="TST").one()
    rec = db.query(Receipt).filter_by(site_id=site.id).one()
    rec.attachment = "abc123key"  # فاتورة مرفوعة جوه البرنامج
    db.commit()
    db.close()

    try:
        run(book(), "TST", "تست")
        raise AssertionError("كان لازم يرفض من غير استبدال")
    except ValueError as e:
        assert "استبدال" in str(e)

    msg = run(book(extra_issue=True, receipts=2), "TST", "تست", replace=True)
    assert "اتمسحت النسخة القديمة" in msg
    db = SessionLocal()
    site = db.query(Site).filter_by(code="TST").one()
    assert db.query(OpeningBalance).filter_by(site_id=site.id).count() == 1
    recs = {r.invoice_no: r for r in db.query(Receipt).filter_by(site_id=site.id)}
    assert set(recs) == {"INV-1", "INV-2"}
    assert recs["INV-1"].attachment == "abc123key" and recs["INV-2"].attachment is None
    assert db.query(Request).filter_by(site_id=site.id, imported=True).count() == 2
    db.close()
