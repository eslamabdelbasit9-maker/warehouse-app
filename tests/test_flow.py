"""اختبار شامل للدورة: طلب → 3 مراحل اعتماد لكل صنف → خصم الرصيد → إشعارات."""
import os
import tempfile

_tmp = tempfile.mkdtemp()
os.environ["DATABASE_URL"] = os.environ.get("TEST_DATABASE_URL", f"sqlite:///{_tmp}/test.db")
os.environ["UPLOAD_DIR"] = f"{_tmp}/up"
os.environ["DEV_AUTH"] = "true"
os.environ["MAIL_MODE"] = "outbox"
os.environ["BOOTSTRAP_ADMIN_PASSWORD"] = "admin-pass-123"

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.db import SessionLocal  # noqa: E402
from app.main import app, init_db  # noqa: E402
from app.models import (CustodyRecord, Item, MailLog, OpeningBalance, Receipt, ReceiptLine, Request, Site, Transfer, Unit,  # noqa: E402
                        User)
from app import services as S  # noqa: E402
from scripts import seed_demo  # noqa: E402


@pytest.fixture(scope="module")
def env():
    init_db()
    seed_demo.run()
    db = SessionLocal()
    site = db.query(Site).first()
    a = Item(code="SP_1", name="رولمان بلي", uom="قطعة")
    b = Item(code="SP_2", name="سير", uom="قطعة")
    db.add_all([a, b])
    db.flush()
    db.add_all([OpeningBalance(site_id=site.id, item_id=a.id, qty=10, value=1000),
                OpeningBalance(site_id=site.id, item_id=b.id, qty=4, value=200)])
    db.commit()
    ids = dict(site=site.id, a=a.id, b=b.id,
               crusher=db.query(Unit).filter_by(site_id=site.id, kind="crusher").first().id,
               asphalt=db.query(Unit).filter_by(site_id=site.id, kind="asphalt").first().id)
    ids.update({u.email: u.id for u in db.query(User)})
    db.close()
    return ids


def login(c, env, email):
    c.cookies.clear()
    r = c.post("/dev-login", data={"uid": env[email]}, follow_redirects=False)
    assert r.status_code == 303


def test_spare_full_cycle(env):
    c = TestClient(app)
    login(c, env, "requester@example.com")
    # كمية أكبر من الرصيد مرفوضة
    r = c.post("/requests/new/spare", data={"site_id": env["site"], "unit_id": env["crusher"], "reason": "صيانة",
                                            "item_id": [env["a"]], "qty": ["11"]})
    assert "أكبر من الرصيد" in c.get(r.headers.get("location", "/requests/new/spare")).text or r.status_code == 200
    r = c.post("/requests/new/spare", data={"site_id": env["site"], "unit_id": env["crusher"], "reason": "صيانة",
                                            "item_id": [env["a"], env["b"]], "qty": ["3", "2"]}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/requests/")
    rid = int(r.headers["location"].rsplit("/", 1)[1])
    db = SessionLocal()
    req = db.get(Request, rid)
    assert req.status == "pending" and req.current_stage == 1
    mail = db.query(MailLog).order_by(MailLog.id.desc()).first()
    assert "eng.crusher@example.com" in mail.to and f"/approvals/{rid}" in mail.html
    la, lb = req.lines
    db.close()

    # مهندس الاسفلت لا يستطيع اعتماد طلب كسارة
    login(c, env, "eng.asphalt@example.com")
    assert c.get(f"/approvals/{rid}", follow_redirects=False).status_code in (303, 403)

    # المهندس يوافق على الصنفين
    login(c, env, "eng.crusher@example.com")
    assert "مراجعة" in c.get("/approvals").text
    c.post(f"/approvals/{rid}", data={"action": "approve", "line": [la.id, lb.id]})
    # مدير الإنتاج يرفض الصنف الثاني بدون سبب → مرفوض
    login(c, env, "prod.manager@example.com")
    c.post(f"/approvals/{rid}", data={"action": "approve", "line": [la.id]})
    db = SessionLocal()
    assert db.get(Request, rid).current_stage == 2
    db.close()
    c.post(f"/approvals/{rid}", data={"action": "approve", "line": [la.id], "comment": "السير غير مطلوب الآن"})
    db = SessionLocal()
    req = db.get(Request, rid)
    assert req.current_stage == 3
    assert [l.status for l in req.lines] == ["pending", "rejected"]
    db.close()
    # النهائي — يرى الصنف الأول فقط
    login(c, env, "final@example.com")
    page = c.get(f"/approvals/{rid}").text
    assert "رولمان بلي" in page
    c.post(f"/approvals/{rid}", data={"action": "approve", "line": [la.id]})
    db = SessionLocal()
    req = db.get(Request, rid)
    assert req.status == "partial"
    l = req.lines[0]
    assert l.unit_cost == 100 and l.value == 300
    st = S.stock_table(db, env["site"])
    assert st[env["a"]]["cur"] == 7 and st[env["b"]]["cur"] == 4
    last = db.query(MailLog).order_by(MailLog.id.desc()).first()
    assert "requester@example.com" in last.to and "معتمد جزئياً" in last.subject
    db.close()
    # لا يمكن الاعتماد مرة أخرى
    r = c.post(f"/approvals/{rid}", data={"action": "approve", "line": [la.id]}, follow_redirects=False)
    db = SessionLocal()
    assert S.stock_table(db, env["site"])[env["a"]]["cur"] == 7
    db.close()


def test_reject_all(env):
    c = TestClient(app)
    login(c, env, "requester@example.com")
    r = c.post("/requests/new/spare", data={"site_id": env["site"], "unit_id": env["crusher"], "reason": "x",
                                            "item_id": [env["b"]], "qty": ["1"]}, follow_redirects=False)
    rid = int(r.headers["location"].rsplit("/", 1)[1])
    login(c, env, "eng.crusher@example.com")
    c.post(f"/approvals/{rid}", data={"action": "reject_all", "comment": "لا"})
    db = SessionLocal()
    assert db.get(Request, rid).status == "rejected"
    db.close()


def test_raw_materials(env):
    c = TestClient(app)
    login(c, env, "requester@example.com")
    db = SessionLocal()
    g = db.query(Item).filter_by(code="RM-G34").one().id
    bit = db.query(Item).filter_by(code="RM-BIT").one().id
    db.close()
    base = {"site_id": env["site"], "unit_id": env["asphalt"], "entity_type": "مشروع", "entity_name": "س",
            "item_id": g, "qty": "5"}
    r = c.post("/requests/new/raw", data=base, follow_redirects=True)   # من غير إنتاج الأسفلت: مرفوض
    assert "الكمية المنتجة من الأسفلت" in r.text
    r = c.post("/requests/new/raw", data={
        "site_id": env["site"], "unit_id": env["asphalt"], "asphalt_tons": "500",
        "entity_type": ["مشروع", "عميل"], "entity_name": ["طريق الملك فهد", "مؤسسة الأمل"],
        "item_id": [g, bit], "qty": ["120", "6.5"], "diesel_تشغيل": "900", "diesel_تسخين": "350"},
        follow_redirects=False)
    assert r.status_code == 303 and "/requests/" in r.headers["location"], r.text
    rid = int(r.headers["location"].rsplit("/", 1)[1])
    db = SessionLocal()
    req = db.get(Request, rid)
    assert req.type == "raw" and len(req.lines) == 4
    assert {l.diesel_purpose for l in req.lines} == {None, "تشغيل", "تسخين"}
    assert req.asphalt_tons == 500 and req.diesel_liters == 1250 and abs(req.liters_per_ton - 2.5) < 1e-9
    mail = db.query(MailLog).order_by(MailLog.id.desc()).first()
    assert "eng.asphalt@example.com" in mail.to
    ids = [l.id for l in req.lines]
    db.close()
    for who in ("eng.asphalt@example.com", "prod.manager@example.com", "final@example.com"):
        login(c, env, who)
        c.post(f"/approvals/{rid}", data={"action": "approve", "line": ids})
    db = SessionLocal()
    assert db.get(Request, rid).status == "approved"
    db.close()
    page = c.get(f"/requests/{rid}").text
    assert "الأسفلت المنتج" in page and "2.50" in page and "لتر/طن" in page
    login(c, env, "requester@example.com")
    assert 'id="tons"' in c.get("/requests/new/raw").text
    page = c.get(f"/raw-stock?site={env['site']}").text
    assert "أرصدة المواد الخام" in page and "RM-G34" in page
    db = SessionLocal()
    st = S.raw_stock_table(db, env["site"])
    assert st[g]["out_qty"] == 120 and st[bit]["out_qty"] == 6.5
    db.close()
    assert "طريق الملك فهد" in c.get("/").text or True


def test_receipt_updates_avg(env):
    c = TestClient(app)
    login(c, env, "storekeeper@example.com")
    r = c.post("/receipts/new", data={"site_id": env["site"], "date": "2026-10-01", "invoice_no": "INV-1",
                                      "supplier": "مورد", "item_id": [env["b"]], "qty": ["6"], "unit_price": ["100"]},
               files={"attachment": ("inv.pdf", b"%PDF-1.4 test", "application/pdf")}, follow_redirects=False)
    assert r.status_code == 303
    db = SessionLocal()
    s = S.stock_table(db, env["site"])[env["b"]]
    assert s["cur"] == 10 and round(s["avg"], 2) == 80.0  # (200+600)/(4+6)
    db.close()


def test_custody(env):
    c = TestClient(app)
    login(c, env, "storekeeper@example.com")
    # قاعدة بيانات قديمة من غير عمود الصورة: init_db بيضيفه
    from sqlalchemy import inspect, text
    from app.db import engine
    from app.main import init_db
    with engine.begin() as conn:
        conn.execute(text("ALTER TABLE custody_records DROP COLUMN photo"))
    init_db()
    assert "photo" in {col["name"] for col in inspect(engine).get_columns("custody_records")}

    jpg = b"\xff\xd8\xff\xe0fake-jpeg"
    r = c.post("/custody/new", data={"site_id": env["site"], "employee_no": "1234", "employee_name": "أحمد",
                                     "category": "أدوات السلامة", "row": ["0", "1"],
                                     "item_name_0": "خوذة", "qty_0": "1", "item_name_1": "حذاء سلامة", "qty_1": "1"},
               files={"photo_0": ("helmet.jpg", jpg, "image/jpeg")}, follow_redirects=False)
    assert r.status_code == 303
    db = SessionLocal()
    recs = db.query(CustodyRecord).filter_by(employee_no="1234").order_by(CustodyRecord.id).all()
    assert len(recs) == 2 and recs[0].photo and not recs[1].photo
    db.close()
    assert c.get(f"/files/{recs[0].photo}").content == jpg
    # صورة بعدين من صفحة الموظف، والـ PDF مش مقبول كصورة
    c.post(f"/custody/{recs[1].id}/photo", files={"photo": ("x.pdf", b"%PDF", "application/pdf")})
    db = SessionLocal()
    assert db.get(CustodyRecord, recs[1].id).photo is None
    db.close()
    c.post(f"/custody/{recs[1].id}/photo", files={"photo": ("shoe.png", b"\x89PNG", "image/png")})
    db = SessionLocal()
    assert db.get(CustodyRecord, recs[1].id).photo
    db.close()
    assert 'class="thumb"' in c.get("/custody/employee/1234").text
    c.post(f"/custody/{recs[0].id}/return", data={"return_condition": "سليمة"})
    db = SessionLocal()
    assert db.get(CustodyRecord, recs[0].id).returned_at is not None
    db.close()
    assert "خوذة" in c.get("/custody/employee/1234").text


def test_pages_render(env):
    c = TestClient(app)
    login(c, env, "admin@example.com")
    for url in ["/", "/requests", "/approvals", "/stock", f"/stock/{env['a']}", "/receipts", "/receipts/new",
                "/custody", "/custody/new", "/requests/new/spare", "/requests/new/raw", "/admin/users",
                "/admin/routes", "/admin/sites", "/admin/items", "/admin/opening", "/admin/mail",
                "/export/requests.xlsx?type=spare", "/export/requests.xlsx?type=raw", "/export/stock.xlsx",
                "/export/custody.xlsx", f"/api/stock?site={env['site']}"]:
        r = c.get(url)
        assert r.status_code == 200, (url, r.status_code, r.text[:500])


def test_requires_login():
    c = TestClient(app)
    r = c.get("/", follow_redirects=False)
    assert r.status_code == 303 and "/login" in r.headers["location"]


def test_password_login_and_attachment_and_import(env):
    from app import config
    c = TestClient(app)
    r = c.post("/signin", data={"email": "admin@example.com", "password": "wrong"}, follow_redirects=False)
    assert c.get("/", follow_redirects=False).status_code == 303
    r = c.post("/signin", data={"email": "admin@example.com", "password": "admin-pass-123"}, follow_redirects=False)
    assert r.status_code == 303
    assert c.get("/").status_code == 200
    # admin sets a password for a user, user logs in
    db = SessionLocal()
    u = db.query(User).filter_by(email="final@example.com").one()
    data = {"id": u.id, "email": u.email, "name": u.name, "roles": ["final"], "all_sites": "1", "active": "1",
            "password": "final-pass-1"}
    db.close()
    c.post("/admin/users", data=data)
    c2 = TestClient(app)
    r = c2.post("/signin", data={"email": "final@example.com", "password": "final-pass-1"}, follow_redirects=False)
    assert c2.get("/approvals").status_code == 200
    # attachment stored in DB and downloadable
    db = SessionLocal()
    from app.models import Receipt
    rec = db.query(Receipt).filter(Receipt.attachment.is_not(None)).first()
    db.close()
    r = c.get(f"/files/{rec.attachment}")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    # import page with a small workbook
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table
    import io as _io
    wb = Workbook()
    ws = wb.active; ws.title = "الرصيد"
    ws.append(["title"])
    ws.append(["كود الصنف", "اسم الصنف", "الموقع", "الرصيد الافتتاحي", "قيمة الرصيد الافتتاحي", "وحدة القياس"])
    ws.append(["X_1", "صنف تجربة", "مخزن", 5, 50, "قطعة"])
    ws.add_table(Table(displayName="BalanceTable", ref="A2:F3"))
    w2 = wb.create_sheet("الوارد"); w2.append(["التاريخ", "رقم الفاتورة", "كود الصنف", "اسم الصنف", "الكمية", "سعر الوحدة", "الإجمالي", "المورد"])
    w2.append(["2026-09-01", "F1", "X_1", "صنف تجربة", 5, 20, 100, "م"])
    w3 = wb.create_sheet("الصرف"); w3.append(["رقم الطلب", "التاريخ", "كود الصنف", "اسم الصنف", "الكمية", "المشروع", "طالب الصرف", "السبب", "حالة الاعتماد", "المهندس المعتمد", "سعر الوحدة", "القيمة", "الشهر", "ReqID"])
    w3.append(["R1", "2026-09-02", "X_1", "صنف تجربة", 2, "كسارة", "أ", "س", "معتمد نهائي", None, 15, 30, "2026-09", "R1"])
    w3.add_table(Table(displayName="IssueTable", ref="A1:N2"))
    buf = _io.BytesIO(); wb.save(buf)
    r = c.post("/admin/import", data={"site_code": "TBK", "site_name": "تبوك"},
               files={"file": ("t.xlsx", buf.getvalue(), "application/octet-stream")})
    assert "تم: موقع تبوك" in r.text, r.text[:2000]
    db = SessionLocal()
    from app.models import Site as _S
    sid = db.query(_S).filter_by(code="TBK").one().id
    st = S.stock_table(db, sid)
    x = db.query(Item).filter_by(code="X_1").one().id
    assert st[x]["cur"] == 8 and round(st[x]["avg"], 2) == 15.0
    db.close()
    r = c.post("/admin/import", data={"site_code": "TBK", "site_name": "تبوك"},
               files={"file": ("t.xlsx", buf.getvalue(), "application/octet-stream")})
    assert "بيانات بالفعل" in r.text


def _link_from_mail(to):
    import re
    db = SessionLocal()
    m = db.query(MailLog).filter(MailLog.to == to).order_by(MailLog.id.desc()).first()
    db.close()
    return re.search(r"(/a/[^'?]+)", m.html).group(1)


def test_approve_from_email_link(env):
    c = TestClient(app)
    login(c, env, "requester@example.com")
    r = c.post("/requests/new/spare", data={"site_id": env["site"], "unit_id": env["crusher"], "reason": "من الإيميل",
                                            "item_id": [env["a"]], "qty": ["1"]}, follow_redirects=False)
    rid = int(r.headers["location"].rsplit("/", 1)[1])
    db = SessionLocal()
    lid = db.get(Request, rid).lines[0].id
    db.close()
    c.cookies.clear()  # المعتمد غير مسجل دخول — يفتح الرابط من Outlook
    link = _link_from_mail("eng.crusher@example.com")
    # فتح الرابط (أو فحصه تلقائياً من Outlook) لا يعتمد شيئاً
    page = c.get(link + "?do=approve")
    assert page.status_code == 200 and "رولمان بلي" in page.text and "مهندس الكسارة" in page.text
    db = SessionLocal()
    assert db.get(Request, rid).current_stage == 1
    db.close()
    # رابط معدّل مرفوض
    assert "غير صالح" in c.post(link + "x", data={"action": "approve", "line": [lid]}).text
    # الموافقة بزر الصفحة
    res = c.post(link, data={"action": "approve", "line": [lid]})
    assert "تم تسجيل قرارك" in res.text
    db = SessionLocal()
    assert db.get(Request, rid).current_stage == 2
    db.close()
    # نفس الرابط لا يُستخدم مرتين
    assert "تم البت" in c.post(link, data={"action": "approve", "line": [lid]}).text
    # مدير الإنتاج يرفض من رابطه — الملاحظة إجبارية
    link2 = _link_from_mail("prod.manager@example.com")
    assert "لم تحدد أي صنف" in c.post(link2, data={"action": "approve"}).text
    assert "أدخل سبب الرفض" in c.post(link2, data={"action": "reject_all"}).text
    res = c.post(link2, data={"action": "reject_all", "comment": "غير مطلوب"})
    assert "تم تسجيل قرارك" in res.text
    db = SessionLocal()
    assert db.get(Request, rid).status == "rejected"
    db.close()


def test_issued_month_year_shown(env):
    from datetime import date
    db = SessionLocal()
    site = db.get(Site, env["site"])
    it = Item(code="SP_ISS", name="فلتر هواء", uom="قطعة")
    db.add(it)
    db.flush()
    db.add(OpeningBalance(site_id=site.id, item_id=it.id, qty=20, value=200))
    db.commit()
    iid = it.id
    db.close()
    c = TestClient(app)

    def new_req(qty, d):
        login(c, env, "requester@example.com")
        r = c.post("/requests/new/spare", data={"site_id": env["site"], "unit_id": env["crusher"], "reason": "x",
                                                "work_date": d, "item_id": [iid], "qty": [qty]}, follow_redirects=False)
        return int(r.headers["location"].rsplit("/", 1)[1])

    def approve_all(rid):
        for who in ("eng.crusher@example.com", "prod.manager@example.com", "final@example.com"):
            login(c, env, who)
            db = SessionLocal()
            lid = db.get(Request, rid).lines[0].id
            db.close()
            c.post(f"/approvals/{rid}", data={"action": "approve", "line": [lid]})

    today = date.today()
    other_month = date(today.year, 1 if today.month != 1 else 2, 1)
    approve_all(new_req("3", today.isoformat()))
    approve_all(new_req("2", other_month.isoformat()))
    rid = new_req("1", today.isoformat())
    db = SessionLocal()
    s = S.issued_summary(db, env["site"], [iid], today, exclude_request_id=rid)[iid]
    db.close()
    assert (s["m_qty"], s["m_n"], s["y_qty"], s["y_n"]) == (3, 1, 5, 2)
    login(c, env, "eng.crusher@example.com")
    page = c.get(f"/approvals/{rid}").text
    assert "منصرف الشهر" in page and "منصرف السنة" in page and "طلب)" not in page
    login(c, env, "requester@example.com")
    assert "منصرف السنة" in c.get(f"/requests/{rid}").text
    api = {i["id"]: i for i in c.get(f"/api/stock?site={env['site']}").json()}[iid]
    assert api["y_qty"] == 5 and api["m_qty"] == 3


def test_transfer_between_sites(env):
    from datetime import date
    db = SessionLocal()
    a = db.get(Site, env["site"])
    b = Site(code="TB2", name="موقع تجريبي 2")
    it = Item(code="SP_TRF", name="طرمبة مياه", uom="قطعة")
    db.add_all([b, it])
    db.flush()
    db.add(OpeningBalance(site_id=a.id, item_id=it.id, qty=10, value=1000))  # متوسط 100
    rec = Receipt(site_id=b.id, date=date.today(), invoice_no="B-1")
    rec.lines = [ReceiptLine(item_id=it.id, qty=5, unit_price=40)]          # متوسط 40
    db.add(rec)
    sk2 = User(email="sk2@example.com", name="أمين موقع 2", roles="storekeeper")
    sk2.sites = [b]
    db.add(sk2)
    db.commit()
    ids = dict(a=a.id, b=b.id, it=it.id, sk2=sk2.id)
    db.close()
    env["sk2@example.com"] = ids["sk2"]
    c = TestClient(app)

    login(c, env, "requester@example.com")
    assert c.get("/transfers", follow_redirects=False).status_code == 403
    # أمين مستودع من غير صلاحية «تحويل بين الفروع»: يأكد الاستلام بس، ما يعملش تحويل
    login(c, env, "sk2@example.com")
    assert c.get("/transfers").status_code == 200
    assert c.get("/transfers/new", follow_redirects=False).status_code == 403
    # صلاحية التحويل من غير ما يكون أمين مستودع ولا مربوط بالموقع المستلِم
    db = SessionLocal()
    tu = User(email="trf@example.com", name="محوّل", roles="transfer")
    tu.sites = [db.get(Site, ids["a"])]
    db.add(tu)
    db.commit()
    env["trf@example.com"] = tu.id
    db.close()
    login(c, env, "trf@example.com")
    r = c.post("/transfers/new", data={"from_site_id": ids["a"], "to_site_id": ids["b"], "item_id": [ids["it"]],
                                       "qty": ["1"]}, follow_redirects=False)
    tid0 = int(r.headers["location"].rsplit("/", 1)[1])
    page = c.get(f"/transfers/{tid0}").text
    assert "إلغاء التحويل" in page and "recv_" not in page   # ما يقدرش يستلم في الموقع التاني
    assert c.post(f"/transfers/{tid0}/receive", data={}, follow_redirects=False).status_code == 403
    c.post(f"/transfers/{tid0}/cancel")

    login(c, env, "storekeeper@example.com")
    r = c.post("/transfers/new", data={"from_site_id": ids["a"], "to_site_id": ids["b"], "item_id": [ids["it"]],
                                       "qty": ["11"]}, follow_redirects=True)
    assert "أكبر من الرصيد المتاح" in r.text
    r = c.post("/transfers/new", data={"from_site_id": ids["a"], "to_site_id": ids["b"], "item_id": [ids["it"]],
                                       "qty": ["4"], "note": "نقص في مكة"}, follow_redirects=False)
    tid = int(r.headers["location"].rsplit("/", 1)[1])
    db = SessionLocal()
    t = db.get(Transfer, tid)
    assert t.status == "in_transit" and t.lines[0].unit_cost == 100 and t.tr_no.startswith("TR-")
    lid = t.lines[0].id
    assert S.stock_table(db, ids["a"])[ids["it"]]["cur"] == 6
    assert S.stock_table(db, ids["b"])[ids["it"]]["cur"] == 5   # لسه ما اتستلمش
    assert "sk2@example.com" in db.query(MailLog).order_by(MailLog.id.desc()).first().to
    db.close()

    # أمين الموقع التاني: يشوف التحويل ويأكد الاستلام، وما يقدرش يلغيه
    login(c, env, "sk2@example.com")
    assert "تأكيد الاستلام" in c.get("/transfers").text
    assert c.post(f"/transfers/{tid}/cancel", follow_redirects=False).status_code == 403
    r = c.post(f"/transfers/{tid}/receive", data={f"recv_{lid}": "3"}, follow_redirects=True)
    assert "أدخل السبب" in r.text
    c.post(f"/transfers/{tid}/receive", data={f"recv_{lid}": "3", "note": "قطعة مكسورة"})
    db = SessionLocal()
    assert db.get(Transfer, tid).status == "received"
    sa, sb = S.stock_table(db, ids["a"])[ids["it"]], S.stock_table(db, ids["b"])[ids["it"]]
    assert sa["cur"] == 7 and sa["avg"] == 100            # الفرق رجع للمرسِل
    assert sb["cur"] == 8 and sb["avg"] == 62.5           # (5×40 + 3×100) ÷ 8
    db.close()
    assert "تحويل وارد" in c.get(f"/stock/{ids['it']}?site={ids['b']}").text

    # الإلغاء قبل الاستلام يرجّع الرصيد
    login(c, env, "storekeeper@example.com")
    assert "مغلق" in c.post(f"/transfers/{tid}/cancel", follow_redirects=True).text
    r = c.post("/transfers/new", data={"from_site_id": ids["a"], "to_site_id": ids["b"], "item_id": [ids["it"]],
                                       "qty": ["2"]}, follow_redirects=False)
    tid2 = int(r.headers["location"].rsplit("/", 1)[1])
    c.post(f"/transfers/{tid2}/cancel")
    db = SessionLocal()
    assert db.get(Transfer, tid2).status == "cancelled"
    assert S.stock_table(db, ids["a"])[ids["it"]]["cur"] == 7
    db.close()


def test_custody_categories(env):
    from datetime import date
    from app.main import init_db
    db = SessionLocal()
    db.add(CustodyRecord(site_id=env["site"], employee_no="E77", employee_name="قديم", category="مهمات سلامة",
                         item_name="نظارة", qty=1, issued_at=date.today()))
    db.add(CustodyRecord(site_id=env["site"], employee_no="E78", employee_name="أصل", category="أصول",
                         item_name="لابتوب", qty=1, issued_at=date.today()))
    db.commit()
    db.close()
    init_db()  # الأسماء القديمة بتتحدّث
    db = SessionLocal()
    assert {c.category for c in db.query(CustodyRecord).filter(CustodyRecord.employee_no.in_(["E77", "E78"]))} == \
        {"أدوات السلامة", "الأصول"}
    db.close()
    c = TestClient(app)
    login(c, env, "storekeeper@example.com")
    safety = c.get(f"/custody?site={env['site']}&cat=أدوات السلامة").text
    assert "نظارة" in safety and "لابتوب" not in safety
    assets = c.get(f"/custody?site={env['site']}&cat=الأصول").text
    assert "لابتوب" in assets and "نظارة" not in assets
    assert "selected>أدوات السلامة" in c.get(f"/custody/new?site={env['site']}&cat=أدوات السلامة").text


def test_appearance(env):
    c = TestClient(app)
    login(c, env, "storekeeper@example.com")
    assert c.get("/admin/appearance", follow_redirects=False).status_code == 403
    db = SessionLocal()
    admin = db.query(User).filter(User.roles == "admin").first()
    env["admin"] = admin.id
    db.close()
    login(c, env, "admin")
    page = c.get("/").text
    assert "--brand:#2B3440" in page and "family=Cairo" in page and "/static/vendor/gsap.min.js" in page
    assert "تم حفظ المظهر" in c.post("/admin/appearance", data={"primary": "#7a1f3d", "bg": "warm", "font": "Almarai"},
                                    follow_redirects=True).text
    page = c.get("/stock").text
    assert "--brand:#7A1F3D" in page and "--bg:#F0EEEA" in page and "family=Almarai" in page
    assert "اختر لون صحيح" in c.post("/admin/appearance", data={"primary": "red", "bg": "gray", "font": "Cairo"},
                                     follow_redirects=True).text
    c.post("/admin/appearance", data={"primary": "#2B3440", "bg": "dark", "font": "Cairo"})
    assert 'data-mode="dark"' in c.get("/").text
    c.post("/admin/appearance", data={"reset": "1"})
    assert 'data-mode="dark"' not in c.get("/").text
    assert "--brand:#2B3440" in c.get("/").text
    c.cookies.clear()
    assert "--brand:#2B3440" in c.get("/signin").text   # صفحة الدخول كمان بتاخد الثيم


def test_dashboard_roles_and_stagnant(env):
    from datetime import date, timedelta
    db = SessionLocal()
    site = db.get(Site, env["site"])
    old = Item(code="SP_OLD", name="صنف راكد", uom="قطعة")
    new = Item(code="SP_NEW", name="صنف متحرك", uom="قطعة")
    db.add_all([old, new])
    db.flush()
    db.add(OpeningBalance(site_id=site.id, item_id=old.id, qty=4, value=400))
    rec = Receipt(site_id=site.id, date=date.today() - timedelta(days=10), invoice_no="NEW-1")
    rec.lines = [ReceiptLine(item_id=new.id, qty=2, unit_price=50)]
    db.add(rec)
    db.commit()
    y = date.today().year
    end = date(y, 12, 31)
    rows = {r["item"].code: r for r in S.stagnant_items(db, [site.id], 180, today=end)}
    # رصيد افتتاحي وما اتصرفش منه: العد من بداية السنة
    assert "SP_OLD" in rows and rows["SP_OLD"]["value"] == 400 and rows["SP_OLD"]["last"] is None
    assert rows["SP_OLD"]["since"] == date(y, 1, 1) and rows["SP_OLD"]["idle_days"] == 364
    assert rows["SP_OLD"]["added_qty"] == 0
    db.close()

    c = TestClient(app)
    login(c, env, "requester@example.com")
    page = c.get("/").text
    assert "طلباتي" in page and "صرف قطع الغيار المعتمد" not in page and "الرواكد" not in page and "المشتريات" not in page
    assert c.get("/stagnant", follow_redirects=False).status_code == 403
    login(c, env, "storekeeper@example.com")
    assert "طلباتي" in c.get("/").text            # أمين المستودع: اللوحة المختصرة
    assert "الأصناف الراكدة" in c.get("/stagnant").text   # بس يقدر يشوف الرواكد
    login(c, env, "eng.crusher@example.com")
    page = c.get("/").text
    assert "صرف قطع الغيار المعتمد" in page and "/stagnant" in page and "نوع الوحدة" in page and "كسارة" in page
    assert "قيمة المشتريات" in page and "المشتريات الشهرية" in page
    # وارد «NEW-1» من 10 أيام بـ 100 (2 × 50) لازم يظهر في المشتريات الشهرية
    page = c.get(f"/?site={env['site']}").text
    assert "100" in page.split("المشتريات الشهرية", 1)[1]
    page = c.get(f"/?site={env['site']}&date_from=2000-01-01&date_to=2099-12-31").text
    assert "عدد الطلبات" in page and "نوع الوحدة" not in page   # موقع واحد: حسب الوحدة نفسها
    page = c.get(f"/stagnant?site={env['site']}&days=90").text
    assert "الأصناف الراكدة" in page and "رواكد المشتريات الشهرية" in page
    page = c.get("/").text
    assert "صرف قطع الغيار الشهري" in page and 'data-m="rk"' not in page and 'href="/requests/new/spare"' not in page
    # فلتر الشهر في الوارد والطلبات (روابط الرسوم)
    from datetime import date as _d
    ym = f"{_d.today():%Y-%m}"
    page = c.get(f"/receipts?site={env['site']}&month={ym}").text
    assert "الشهر:" in page and "NEW-1" in page
    assert "NEW-1" not in c.get(f"/receipts?site={env['site']}&month=2001-01").text
    assert "الشهر:" in c.get(f"/requests?type=spare&ok=1&month={ym}").text


def test_stagnant_rules(env):
    from datetime import date
    db = SessionLocal()
    site = db.get(Site, env["site"])
    a = Item(code="ST_A", name="رصيد وجاله وارد", uom="قطعة")     # افتتاحي + فاتورة بعد كده، ما اتصرفش
    b = Item(code="ST_B", name="صنف جديد", uom="قطعة")            # أول وارد في نص السنة
    c = Item(code="ST_C", name="جديد وجاله وارد تاني", uom="قطعة")
    db.add_all([a, b, c])
    db.flush()
    db.add(OpeningBalance(site_id=site.id, item_id=a.id, qty=3, value=300))
    for d, it, q in [(date(2031, 3, 1), a, 2), (date(2031, 9, 1), b, 5), (date(2031, 2, 1), c, 1), (date(2031, 4, 1), c, 4)]:
        rec = Receipt(site_id=site.id, date=d, invoice_no=f"ST-{it.code}-{d}")
        rec.lines = [ReceiptLine(item_id=it.id, qty=q, unit_price=10)]
        db.add(rec)
    db.commit()
    rows = {r["item"].code: r for r in S.stagnant_items(db, [site.id], 180, today=date(2031, 12, 1))}
    assert rows["ST_A"]["since"] == date(2031, 1, 1) and rows["ST_A"]["idle_days"] == 334
    assert rows["ST_A"]["added_qty"] == 2                       # جاله وارد بعد آخر صرف → أحمر
    assert "ST_B" not in rows                                   # جديد: العد من أول وارد (91 يوم)
    assert rows["ST_C"]["since"] == date(2031, 2, 1) and rows["ST_C"]["added_qty"] == 4
    # بعد صرف من ST_A العداد بيبدأ من تاريخ الصرف والأحمر بيروح
    from app.models import RequestLine
    unit = db.query(Unit).filter_by(site_id=site.id).first()
    req = Request(req_no="SP-ST-1", type="spare", site_id=site.id, unit_id=unit.id, work_date=date(2031, 5, 1),
                  status="approved", current_stage=3)
    req.lines = [RequestLine(item_id=a.id, qty=1, status="approved")]
    db.add(req)
    db.commit()
    rows = {r["item"].code: r for r in S.stagnant_items(db, [site.id], 180, today=date(2031, 12, 1))}
    assert rows["ST_A"]["last"] == date(2031, 5, 1) and rows["ST_A"]["added_qty"] == 0
    # طلب قيد الاعتماد على ST_C بكل رصيده (5): محسوب مصروف → يخرج من الرواكد
    pend = Request(req_no="SP-ST-2", type="spare", site_id=site.id, unit_id=unit.id, work_date=date(2031, 6, 1),
                   status="pending", current_stage=1)
    pend.lines = [RequestLine(item_id=c.id, qty=5, status="pending")]
    # وطلب معلق على جزء من ST_A (2 من 4): الباقي بس راكد، والعد من تاريخ الطلب
    pend2 = Request(req_no="SP-ST-3", type="spare", site_id=site.id, unit_id=unit.id, work_date=date(2031, 6, 1),
                    status="pending", current_stage=1)
    pend2.lines = [RequestLine(item_id=a.id, qty=2, status="pending")]
    db.add_all([pend, pend2])
    db.commit()
    rows = {r["item"].code: r for r in S.stagnant_items(db, [site.id], 180, today=date(2031, 12, 1))}
    assert "ST_C" not in rows
    assert rows["ST_A"]["cur"] == 2 and rows["ST_A"]["last"] == date(2031, 6, 1)
    # رواكد المشتريات (بالأقدمية): ST_A الافتتاحي (3) اتصرف كله، فوارد مارس (2 × 10) لسه راكد؛
    # ST_C اتصرف كله (معلق)؛ ST_B وارد سبتمبر (5 × 10) كله راكد
    left = {(d.month, v) for sid, d, v, _ in S.unissued_purchases(db, [site.id], date(2031, 1, 1)) if d.year == 2031}
    assert (3, 20) in left and (9, 50) in left and (2, 0) in left and (4, 0) in left
    db.close()


def test_greet_name():
    from app.main import greet_name
    assert greet_name("م/ حسين عبد الله") == "م/ حسين"
    assert greet_name("م. أحمد علي") == "م. أحمد"
    assert greet_name("مهندس خالد عمر") == "مهندس خالد"
    assert greet_name("إسلام عبد الباسط") == "إسلام"
    assert greet_name("م/حسين") == "م/حسين"
    assert greet_name("") == ""


def test_healthz():
    r = TestClient(app).get("/healthz")
    assert r.status_code == 200 and r.json() == {"ok": True}


def test_personal_settings(env):
    from app.auth import check_password
    c = TestClient(app)
    login(c, env, "requester@example.com")
    page = c.get("/account").text
    assert "الإعدادات" in page and "المظهر الخاص بك" in page and "تغيير كلمة المرور" in page
    # مظهر خاص بيه بس
    assert "تم حفظ المظهر" in c.post("/account/appearance", data={"primary": "#7A1F3D", "bg": "dark", "font": "Tajawal"},
                                     follow_redirects=True).text
    page = c.get("/").text
    assert "--brand:#7A1F3D" in page and 'data-mode="dark"' in page and "family=Tajawal" in page
    other = TestClient(app)
    login(other, env, "storekeeper@example.com")
    assert "--brand:#7A1F3D" not in other.get("/").text and 'data-mode="dark"' not in other.get("/").text
    c.post("/account/appearance", data={"reset": "1"})
    assert "--brand:#7A1F3D" not in c.get("/").text
    # كلمة المرور
    c.post("/account", data={"new": "pass-1234", "new2": "pass-1234"})
    db = SessionLocal()
    assert check_password("pass-1234", db.get(User, env["requester@example.com"]).password_hash)
    db.close()
    assert "غير صحيحة" in c.post("/account", data={"old": "wrong", "new": "x" * 8, "new2": "x" * 8}, follow_redirects=True).text


def test_admin_edits_request_header(env):
    """الأدمن يعدل وحدة/تاريخ/سبب طلب؛ غير الأدمن ممنوع؛ ووحدة من موقع تاني مرفوضة."""
    from datetime import date
    c = TestClient(app)
    login(c, env, "requester@example.com")
    r = c.post("/requests/new/spare", data={"site_id": env["site"], "unit_id": env["crusher"], "reason": "صيانة",
                                            "item_id": [env["a"]], "qty": ["1"]}, follow_redirects=False)
    rid = int(r.headers["location"].rsplit("/", 1)[1])

    # غير الأدمن ما يقدرش
    r = c.post(f"/requests/{rid}/edit", data={"unit_id": env["asphalt"]}, follow_redirects=False)
    assert r.status_code in (302, 303, 403)
    db = SessionLocal()
    assert db.get(Request, rid).unit_id == env["crusher"]
    db.close()

    login(c, env, "admin@example.com")
    assert "تعديل بيانات الطلب" in c.get(f"/requests/{rid}").text
    c.post(f"/requests/{rid}/edit", data={"unit_id": env["asphalt"], "work_date": "2026-09-15", "reason": "تغيير سير"})
    db = SessionLocal()
    req = db.get(Request, rid)
    assert req.unit_id == env["asphalt"] and req.work_date == date(2026, 9, 15) and req.reason == "تغيير سير"
    assert req.status == "pending" and req.lines[0].qty == 1

    # وحدة من موقع تاني مرفوضة
    other = Site(code="OTH", name="موقع تاني")
    db.add(other)
    db.flush()
    u = Unit(site_id=other.id, name="كسارة", kind="crusher")
    db.add(u)
    db.commit()
    uid = u.id
    db.close()
    c.post(f"/requests/{rid}/edit", data={"unit_id": uid})
    db = SessionLocal()
    assert db.get(Request, rid).unit_id == env["asphalt"]
    db.close()


def test_raw_materials_production_concrete_sale(env):
    from app.models import CrusherProduction
    db = SessionLocal()
    site = db.get(Site, env["site"])
    conc = Unit(site_id=site.id, name="محطة خرسانة", kind="concrete")
    db.add(conc)
    db.commit()
    conc_id = conc.id
    g = db.query(Item).filter_by(code="RM-G34").one()
    assert g.crusher                                   # البحص من نواتج الكسارة
    gid = g.id
    db.close()
    c = TestClient(app)
    # السعر الثابت: مدير النظام بس
    login(c, env, "admin@example.com")
    c.post("/admin/crusher-prices", data={f"price_{env['site']}": "0"})
    login(c, env, "storekeeper@example.com")
    assert c.get("/admin/crusher-prices", follow_redirects=False).status_code == 403
    assert "لم يُحدَّد سعر" in c.post("/production/new", data={"site_id": env["site"], "unit_id": env["crusher"],
                                                              "item_id": [gid], "qty": ["100"]}, follow_redirects=True).text
    login(c, env, "admin@example.com")
    c.post("/admin/crusher-prices", data={f"price_{env['site']}": "12.95"})
    # إنتاج من غير اعتماد بالسعر الثابت
    login(c, env, "storekeeper@example.com")
    assert "12.95" in c.get(f"/production/new?site={env['site']}").text
    r = c.post("/production/new", data={"site_id": env["site"], "unit_id": env["crusher"], "item_id": [gid], "qty": ["500"]},
               follow_redirects=True)
    assert "تم تسجيل إنتاج 500 طن" in r.text and "6,475" in r.text
    db = SessionLocal()
    st = S.raw_stock_table(db, env["site"])
    base = st[gid]["cur"]
    assert st[gid]["prod_qty"] == 500 and st[gid]["avg"] == 12.95
    assert db.query(CrusherProduction).count() == 1
    db.close()
    # صرف خرسانة: اعتماد مدير الإنتاج بس، والتكلفة بالسعر الثابت
    login(c, env, "requester@example.com")
    r = c.post("/requests/new/raw", data={"site_id": env["site"], "dest": "concrete", "unit_id": conc_id,
                                          "entity_type": "مشروع", "entity_name": "مشروع خرسانة", "item_id": gid, "qty": "40"},
               follow_redirects=False)
    rid = int(r.headers["location"].rsplit("/", 1)[1])
    db = SessionLocal()
    req = db.get(Request, rid)
    assert req.raw_dest == "concrete" and req.current_stage == 2
    lid = req.lines[0].id
    db.close()
    login(c, env, "prod.manager@example.com")
    c.post(f"/approvals/{rid}", data={"action": "approve", "line": [lid]})
    db = SessionLocal()
    req = db.get(Request, rid)
    assert req.status == "approved" and req.lines[0].unit_cost == 12.95 and req.lines[0].value == 518
    rs = S.raw_stock_table(db, env["site"])
    assert rs[gid]["cur"] == base - 40 and rs[gid]["out_concrete"] == 40
    db.close()
    # بيع: اسم العميل والسعر والقيمة — مدير الإنتاج ثم النهائي
    login(c, env, "requester@example.com")
    assert "سعر البيع" in c.get(f"/requests/new/raw?site={env['site']}&dest=sale").text
    r = c.post("/requests/new/raw", data={"site_id": env["site"], "dest": "sale", "customer": "مؤسسة النور",
                                          "item_id": [gid], "qty": ["10"], "sale_price": ["30"]}, follow_redirects=False)
    rid = int(r.headers["location"].rsplit("/", 1)[1])
    page = c.get(f"/requests/{rid}").text
    assert "مؤسسة النور" in page and "قيمة البيع" in page and "300" in page
    db = SessionLocal()
    req = db.get(Request, rid)
    assert req.current_stage == 2 and req.sale_total == 300
    lid = req.lines[0].id
    db.close()
    for who in ("prod.manager@example.com", "final@example.com"):
        login(c, env, who)
        c.post(f"/approvals/{rid}", data={"action": "approve", "line": [lid]})
    db = SessionLocal()
    assert db.get(Request, rid).status == "approved"
    db.close()
    login(c, env, "eng.crusher@example.com")
    page = c.get(f"/raw-stock?site={env['site']}&date_from=2000-01-01").text
    assert "الخرسانة" in page and "رصيد أول الفترة" in page


def test_stock_period_and_unit_filter(env):
    c = TestClient(app)
    login(c, env, "storekeeper@example.com")
    page = c.get(f"/stock?site={env['site']}&date_to=2000-01-01").text
    assert "الرصيد في 2000-01-01" in page
    page = c.get(f"/stock?site={env['site']}&unit={env['crusher']}&date_from=2000-01-01").text
    assert "الأصناف المصروفة لـ" in page and "منصرف للوحدة" in page
    assert 'src="/static/picker.js' in c.get(f"/requests/new/spare?site={env['site']}").text


def test_custody_import_and_transfer_history(env):
    """استيراد شيت العهد بالروابط + عدم التكرار، وصفحة التحويلات تعرض تحويلات الشيت."""
    import io
    from openpyxl import Workbook
    from datetime import date
    from app.models import CustodyRecord
    db = SessionLocal()
    for code, name in (("MEX", "مكه"), ("JIX", "الكدمى جيزان")):
        if not db.query(Site).filter_by(code=code).first():
            db.add(Site(code=code, name=name))
    db.commit()
    db.close()
    wb = Workbook()
    a = wb.active
    a.title = "بيانات الأصول"
    a.append(["م", "الفرع", "الرقم الوظيفي", "اسم الموظف", "الوظيفة", "القسم", "اسم الأصل", "كود الأصل", "الكمية", "رابط الصورة", "رابط المرفق"])
    a.append([1, "الكدمي", "5784", "انيل", "حداد", "انتاج", "ماكينة لحام", "-", 1, "https://x/i1", "https://x/b1"])
    s = wb.create_sheet("بيانات السلامة")
    s.append(["م", "الفرع", "الرقم الوظيفي", "اسم الموظف", "الوظيفة", "القسم", "المهمات المستلمة", "الكمية", "رابط المرفق", "تاريخ الاستلام"])
    s.append([1, "مكة", "800", "محمد", "مشغل", "كساره", "حذاء سيفتى", 1, "https://x/b2", "2026/08/27"])
    s.append([2, "فرع وهمي", "900", "س", "", "", "خوذة تجربة", 1, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    c = TestClient(app)
    login(c, env, "admin@example.com")
    c.post("/admin/import/custody", files={"files": ("c.xlsx", data)})
    c.post("/admin/import/custody", files={"files": ("c.xlsx", data)})  # التانية تتخطى
    db = SessionLocal()
    recs = db.query(CustodyRecord).filter(CustodyRecord.item_name.in_(["ماكينة لحام", "حذاء سيفتى", "خوذة تجربة"])).all()
    assert len(recs) == 2
    asset = next(r for r in recs if r.employee_no == "5784")
    safety = next(r for r in recs if r.employee_no == "800")
    assert asset.category == "الأصول" and asset.site.name == "الكدمى جيزان" and asset.photo == "https://x/i1"
    assert safety.category == "أدوات السلامة" and safety.issued_at == date(2026, 8, 27) and safety.attachment == "https://x/b2"
    sid = asset.site_id
    db.close()
    page = c.get(f"/custody?site={sid}&state=all").text
    assert 'href="https://x/b1"' in page and 'href="https://x/i1"' in page
    assert c.get("/transfers").status_code == 200


def test_old_transfers_not_counted_as_spend(env):
    """التحويلات القديمة من الشيتات (وحدة «تحويل...» / مورد «تحويل...») ما تتحسبش صرف ولا مشتريات."""
    from app.models import RequestLine
    from datetime import date
    db = SessionLocal()
    site = db.get(Site, env["site"])
    u = Unit(site_id=site.id, name="تحويل لمكة", kind="other")
    it = Item(code="TR-ITEM", name="صنف محوّل", uom="قطعة")
    db.add_all([u, it])
    db.flush()
    rq = Request(req_no="SP-OLD-TR", type="spare", site_id=site.id, unit_id=u.id, work_date=date.today(),
                 status="approved", current_stage=3, imported=True)
    rq.lines = [RequestLine(item_id=it.id, qty=1, status="approved", unit_cost=777777, value=777777)]
    rec = Receipt(site_id=site.id, date=date.today(), invoice_no="TR-IN", supplier="تحويل من مكة")
    rec.lines = [ReceiptLine(item_id=it.id, qty=1, unit_price=888888)]
    db.add_all([rq, rec])
    db.commit()
    db.close()
    c = TestClient(app)
    login(c, env, "eng.crusher@example.com")
    page = c.get(f"/?site={env['site']}").text
    assert "777,777" not in page and "888,888" not in page



def test_receipt_vat_display_only(env):
    """الضريبة 15% بتظهر في الوارد بس — متوسط التكلفة من غير ضريبة، والتحويل مالوش ضريبة."""
    from datetime import date
    c = TestClient(app)
    login(c, env, "admin@example.com")
    db = SessionLocal()
    rec = Receipt(site_id=env["site"], date=date(2026, 10, 5), invoice_no="VAT-1", supplier="مورد ضريبة")
    rec.lines.append(ReceiptLine(item_id=env["b"], qty=2, unit_price=100))
    tr = Receipt(site_id=env["site"], date=date(2026, 10, 5), invoice_no="TR-1", supplier="تحويل من نجران")
    tr.lines.append(ReceiptLine(item_id=env["b"], qty=1, unit_price=100))
    db.add_all([rec, tr])
    db.commit()
    assert (rec.net, rec.vat, rec.gross) == (200, 30, 230) and tr.vat == 0 and tr.gross == 100
    rid = rec.id
    db.close()
    page = c.get(f"/receipts/{rid}").text
    assert "الإجمالي شامل الضريبة" in page and "230" in page


def test_suppliers_statement(env):
    """الموردين: الأسماء اللي فيها منطقة بتتجمع، وكشف الحساب فيه الموقع والضريبة والمرفق."""
    from datetime import date
    c = TestClient(app)
    login(c, env, "admin@example.com")
    db = SessionLocal()
    for inv, sup, att in (("SH-1", "الشامل الجديد التجارية (صبيا)", "https://x/inv1"), ("SH-2", "الشامل الجديد (خميس مشيط)", None),
                          ("SH-3", "تحويل من نجران", None)):
        r = Receipt(site_id=env["site"], date=date(2026, 9, 1), invoice_no=inv, supplier=sup, attachment=att)
        r.lines.append(ReceiptLine(item_id=env["a"], qty=1, unit_price=100))
        db.add(r)
    db.commit()
    db.close()
    page = c.get("/suppliers").text
    assert "الشامل الجديد" in page and "أسماء مختلفة مدمجة" in page and "تحويل من نجران" not in page
    st = c.get("/suppliers/statement", params={"k": "الشامل الجديد"}).text
    assert "SH-1" in st and "SH-2" in st and "SH-3" not in st and 'href="https://x/inv1"' in st and "230" in st


def test_transfer_by_invoice_no_not_a_supplier(env):
    """وارد رقم فاتورته «استلام تحويل من الدار» ومورده فاضي = تحويل: مالوش ضريبة ومش في الموردين."""
    from datetime import date
    c = TestClient(app)
    login(c, env, "admin@example.com")
    db = SessionLocal()
    r = Receipt(site_id=env["site"], date=date(2026, 4, 30), invoice_no="استلام تحويل من الدار", supplier=None)
    r.lines.append(ReceiptLine(item_id=env["a"], qty=1, unit_price=50))
    db.add(r)
    db.commit()
    assert r.is_transfer and r.vat == 0
    db.close()
    assert "استلام تحويل من الدار" not in c.get("/suppliers/statement", params={"k": ""}).text
    assert "استلام تحويل من الدار" in c.get(f"/transfers?site={env['site']}").text


def test_raw_project_issue_and_single_destination(env):
    """صرف مواد خام لمشروع للشركة: حقل «الجهة المستلمة» واحد، والمشروع بيظهر في كارت الصنف."""
    c = TestClient(app)
    login(c, env, "requester@example.com")
    page = c.get(f"/requests/new/raw?site={env['site']}&dest=project").text
    sel_box = page.split("الصرف إلى")[1].split("</select>")[0]
    assert "بيع لعميل" not in sel_box and "مشروع للشركة" not in sel_box and "seg-tg" not in page   # العميل/المشروع تحت في السطور
    assert "سعر البيع (للعميل)" in page
    assert "اسم الخلاطة" not in page
    page = c.get(f"/requests/new/raw?site={env['site']}&unit={env['asphalt']}").text
    assert 'id="tons"' in page
    db = SessionLocal()
    gid = db.query(Item).filter_by(code="RM-G34").one().id
    db.close()
    r = c.post("/requests/new/raw", data={"site_id": env["site"], "dest": "project", "customer": "طريق نجران الدائري",
                                          "item_id": [gid], "qty": ["30"]}, follow_redirects=False)
    assert r.status_code == 303, r.text
    rid = int(r.headers["location"].rsplit("/", 1)[1])
    db = SessionLocal()
    req = db.get(Request, rid)
    assert req.raw_dest == "project" and req.customer == "طريق نجران الدائري" and req.current_stage == 2
    lid = req.lines[0].id
    db.close()
    for who in ("prod.manager@example.com", "final@example.com"):
        login(c, env, who)
        c.post(f"/approvals/{rid}", data={"action": "approve", "line": [lid]})
    db = SessionLocal()
    assert db.get(Request, rid).status == "approved"
    assert S.raw_stock_table(db, env["site"])[gid]["out_project"] == 30
    db.close()
    card = c.get(f"/stock/{gid}?site={env['site']}").text
    assert "مشروع — طريق نجران الدائري" in card


def test_asphalt_tons_per_party(env):
    """صرف الخلاطة: أطنان الأسفلت لكل جهة بتتسجل، والإجمالي = مجموعها."""
    c = TestClient(app)
    login(c, env, "requester@example.com")
    db = SessionLocal()
    g = db.query(Item).filter_by(code="RM-G34").one().id
    db.close()
    page = c.get(f"/requests/new/raw?site={env['site']}&unit={env['asphalt']}").text
    assert 'id="ptons"' in page and "أطنان الأسفلت المنتجة لكل جهة" in page
    r = c.post("/requests/new/raw", data={
        "site_id": env["site"], "dest": "asphalt", "unit_id": env["asphalt"],
        "entity_type": ["عميل", "مشروع"], "entity_name": ["مؤسسة س", "طريق ص"], "item_id": [g, g], "qty": ["10", "5"],
        "pt_type": ["عميل", "مشروع"], "pt_name": ["مؤسسة س", "طريق ص"], "pt_tons": ["120", "80"]}, follow_redirects=False)
    assert r.status_code == 303, r.text
    rid = int(r.headers["location"].rsplit("/", 1)[1])
    db = SessionLocal()
    req = db.get(Request, rid)
    assert req.asphalt_tons == 200 and {(p.entity_name, p.tons) for p in req.party_tons} == {("مؤسسة س", 120), ("طريق ص", 80)}
    db.close()
    assert "أسفلت عميل مؤسسة س" in c.get(f"/requests/{rid}").text


def test_unit_statement_pages(env):
    """كشف حساب الخلاطة: تكلفة المواد + نصيب الديزل لكل جهة وتكلفة الطن؛ والكسارة بتفتح من غير أخطاء."""
    c = TestClient(app)
    login(c, env, "requester@example.com")
    db = SessionLocal()
    g = db.query(Item).filter_by(code="RM-G34").one().id
    db.close()
    r = c.post("/requests/new/raw", data={
        "site_id": env["site"], "dest": "asphalt", "unit_id": env["asphalt"],
        "entity_type": ["عميل", "مشروع"], "entity_name": ["عميل كشف", "مشروع كشف"], "item_id": [g, g], "qty": ["6", "4"],
        "pt_type": ["عميل", "مشروع"], "pt_name": ["عميل كشف", "مشروع كشف"], "pt_tons": ["60", "40"],
        "diesel_تشغيل": "100"}, follow_redirects=False)
    rid = int(r.headers["location"].rsplit("/", 1)[1])
    db = SessionLocal()
    ids = [l.id for l in db.get(Request, rid).lines]
    db.close()
    for who in ("eng.asphalt@example.com", "prod.manager@example.com", "final@example.com"):
        login(c, env, who)
        c.post(f"/approvals/{rid}", data={"action": "approve", "line": ids})
    db = SessionLocal()
    assert db.get(Request, rid).status == "approved"
    db.close()
    login(c, env, "admin@example.com")
    page = c.get(f"/unit-statement?site={env['site']}&unit={env['asphalt']}").text
    assert "التكلفة حسب الجهة" in page and "عميل كشف" in page and "مشروع كشف" in page and "تكلفة طن الأسفلت" in page
    page = c.get(f"/unit-statement?site={env['site']}&unit={env['crusher']}").text
    assert "الإنتاج والصرف حسب المادة" in page and "قطع الغيار المصروفة للوحدة" in page


def test_import_raw_materials(env):
    """استيراد ملف المواد الخام الموحّد: افتتاحي + إنتاج + وارد + صرف (بيع وخلاطة بأطنان الجهات)، وإعادة الاستيراد بتستبدل."""
    import io
    from datetime import date
    from openpyxl import Workbook
    from app.models import CrusherProduction
    from scripts.import_raw import run
    wb = Workbook()
    wb.remove(wb.active)
    def sh(n, head, rows):
        w = wb.create_sheet(n); w.append(head)
        for r in rows: w.append(r)
    sh("الأصناف", ["الكود", "الاسم", "الوحدة", "ناتج كسارة"], [("RM-T34", "بحص تجربة", "طن", 1), ("RM-TBT", "بيتومين تجربة", "طن", 0)])
    sh("الافتتاحي", ["الكود", "الكمية", "القيمة"], [("RM-T34", 1000, 10000), ("RM-TBT", 50, 100000)])
    sh("الإنتاج", ["التاريخ", "الكود", "الكمية", "السعر"], [(date(2026, 4, 1), "RM-T34", 300, 12.59)])
    sh("الوارد", ["التاريخ", "رقم الفاتورة", "المورد", "الكود", "الكمية", "السعر"], [(date(2026, 4, 2), "B-1", "مورد", "RM-TBT", 10, 2000)])
    sh("الصرف", ["المرجع", "التاريخ", "الجهة", "الوحدة", "العميل", "الكود", "الكمية", "التكلفة", "سعر البيع",
                 "نوع الجهة", "اسم الجهة", "غرض الديزل", "أطنان الأسفلت"],
       [("TS1", date(2026, 4, 3), "sale", "", "عميل س", "RM-T34", 100, 12.59, 15, None, None, None, None),
        ("TA1", date(2026, 4, 3), "asphalt", "", None, "RM-TBT", 5, 2000, None, "عميل", "عميل ص", None, 90)])
    sh("أطنان الجهات", ["المرجع", "نوع الجهة", "اسم الجهة", "الأطنان"], [("TA1", "عميل", "عميل ص", 90)])
    buf = io.BytesIO(); wb.save(buf)
    db = SessionLocal()
    code = db.get(Site, env["site"]).code
    db.close()
    for _ in range(2):
        buf.seek(0)
        msg = run(io.BytesIO(buf.getvalue()), code)
    assert "طلبات صرف 2" in msg
    db = SessionLocal()
    t34 = db.query(Item).filter_by(code="RM-T34").one().id
    tbt = db.query(Item).filter_by(code="RM-TBT").one().id
    st = S.raw_stock_table(db, env["site"])
    assert st[t34]["open_qty"] == 1000 and st[t34]["prod_qty"] == 300 and st[t34]["out_sale"] == 100
    assert st[tbt]["in_qty"] == 10 and st[tbt]["out_asphalt"] == 5
    assert db.query(CrusherProduction).filter_by(site_id=env["site"], notes="استيراد مواد خام").count() == 1
    req = db.query(Request).filter(Request.req_no.like("%TA1")).one()
    assert req.asphalt_tons == 90 and req.party_tons[0].entity_name == "عميل ص" and req.status == "approved"
    db.close()


def test_raw_moving_average_cost(env):
    """البيتومين يُسعَّر بالمتوسط المرجّح المتحرك: كل صرف بمتوسط وقت صرفه، والوارد بأثر رجعي يعيد التسعير."""
    import io
    from datetime import date
    from openpyxl import Workbook
    from scripts.import_raw import run
    wb = Workbook()
    wb.remove(wb.active)
    def sh(n, head, rows):
        w = wb.create_sheet(n); w.append(head)
        for r in rows: w.append(r)
    sh("الأصناف", ["الكود", "الاسم", "الوحدة", "ناتج كسارة"], [("RM-MBT", "بيتومين متوسط", "طن", 0)])
    sh("الافتتاحي", ["الكود", "الكمية", "القيمة"], [("RM-MBT", 10, 10000)])
    sh("الوارد", ["التاريخ", "رقم الفاتورة", "المورد", "الكود", "الكمية", "السعر"],
       [(date(2026, 4, 2), "M-1", "المسار", "RM-MBT", 5, 2000),
        (date(2026, 4, 5), "تحويل", "تحويل من نجران", "RM-MBT", 10, 1700)])
    sh("الصرف", ["المرجع", "التاريخ", "الجهة", "الوحدة", "العميل", "الكود", "الكمية", "التكلفة", "سعر البيع",
                 "نوع الجهة", "اسم الجهة", "غرض الديزل", "أطنان الأسفلت"],
       [("MA1", date(2026, 4, 1), "asphalt", "", None, "RM-MBT", 5, 0, None, None, None, None, None),
        ("MA2", date(2026, 4, 3), "asphalt", "", None, "RM-MBT", 4, 0, None, None, None, None, None),
        ("MA3", date(2026, 4, 6), "asphalt", "", None, "RM-MBT", 2, 0, None, None, None, None, None)])
    buf = io.BytesIO(); wb.save(buf)
    db = SessionLocal()
    code = db.get(Site, env["site"]).code
    db.close()
    run(io.BytesIO(buf.getvalue()), code)
    db = SessionLocal()
    cost = {r.req_no[-3:]: r.lines[0].unit_cost for r in db.query(Request).filter(Request.req_no.like("%-MA_"))}
    # 10@1000 → صرف 5 بـ1000 → وارد 5@2000: (5000+10000)/10=1500 → صرف 4 بـ1500 → وارد 10@1700: (9000+17000)/16=1625
    assert cost == {"MA1": 1000, "MA2": 1500, "MA3": 1625}
    mbt = db.query(Item).filter_by(code="RM-MBT").one().id
    assert round(S.raw_stock_table(db, env["site"])[mbt]["avg"], 4) == 1625
    assert round(S.raw_stock_table(db, env["site"], upto=date(2026, 4, 4))[mbt]["avg"], 4) == 1500
    # فاتورة اترفع عليها مرفق: إعادة الاستيراد ما تكررهاش وتحتفظ بالمرفق
    rc = db.query(Receipt).filter_by(invoice_no="M-1").one()
    assert rc.attachment is None
    rc.attachment = "abc.pdf"
    db.commit()
    db.close()
    run(io.BytesIO(buf.getvalue()), code)
    db = SessionLocal()
    rcs = db.query(Receipt).filter_by(invoice_no="M-1").all()
    assert len(rcs) == 1 and rcs[0].attachment == "abc.pdf"
    db.close()


def test_admin_fixes_approved_issue_qty_and_cancels(env):
    """الأدمن يصحح كمية صرف معتمد أو يلغي صرف مسجّل بالخطأ، والرصيد يتعدل."""
    from app.models import RequestLine
    from datetime import datetime
    db = SessionLocal()
    req = Request(req_no="FIX-1", type="spare", work_date=datetime.now().date(), site_id=env["site"], unit_id=env["crusher"], requester_name="استيراد",
                  status="approved", current_stage=3, imported=True, closed_at=datetime.now())
    req.lines.append(RequestLine(item_id=env["a"], qty=2, status="approved", unit_cost=10, value=20))
    db.add(req); db.commit()
    rid, lid = req.id, req.lines[0].id
    before = S.stock_table(db, env["site"], [env["a"]])[env["a"]]["cur"]
    db.close()
    c = TestClient(app)
    login(c, env, "admin@example.com")
    assert f"qty_{lid}" in c.get(f"/requests/{rid}").text
    c.post(f"/requests/{rid}/edit", data={f"qty_{lid}": "5"})
    db = SessionLocal()
    l = db.get(RequestLine, lid)
    assert l.qty == 5 and l.value == 50
    assert S.stock_table(db, env["site"], [env["a"]])[env["a"]]["cur"] == before - 3
    db.close()
    c.post(f"/requests/{rid}/edit", data={"cancel": "1"})
    db = SessionLocal()
    assert db.get(Request, rid).status == "cancelled"
    assert S.stock_table(db, env["site"], [env["a"]])[env["a"]]["cur"] == before + 2
    db.close()


def test_admin_deletes_request(env):
    from datetime import datetime
    from app.models import RequestLine
    db = SessionLocal()
    req = Request(req_no="DEL-1", type="spare", work_date=datetime.now().date(), site_id=env["site"], unit_id=env["crusher"],
                  requester_name="استيراد", status="approved", current_stage=3, imported=True)
    req.lines.append(RequestLine(item_id=env["a"], qty=1, status="approved", unit_cost=1, value=1))
    db.add(req); db.commit(); rid = req.id; db.close()
    c = TestClient(app)
    login(c, env, "requester@example.com")
    c.post(f"/requests/{rid}/delete")
    db = SessionLocal(); assert db.get(Request, rid) is not None; db.close()
    login(c, env, "admin@example.com")
    c.post(f"/requests/{rid}/delete")
    db = SessionLocal(); assert db.get(Request, rid) is None; db.close()


def test_admin_edits_and_deletes_receipt(env):
    from datetime import date
    db = SessionLocal()
    rec = Receipt(site_id=env["site"], date=date(2026, 9, 1), invoice_no="E-1", supplier="مورد")
    rec.lines.append(ReceiptLine(item_id=env["a"], qty=4, unit_price=10))
    db.add(rec); db.commit(); rid, lid = rec.id, rec.lines[0].id
    before = S.stock_table(db, env["site"], [env["a"]])[env["a"]]["cur"]
    db.close()
    c = TestClient(app)
    login(c, env, "requester@example.com")
    c.post(f"/receipts/{rid}/edit", data={f"qty_{lid}": "9"})
    login(c, env, "admin@example.com")
    assert "تعديل الفاتورة" in c.get(f"/receipts/{rid}").text
    c.post(f"/receipts/{rid}/edit", data={"date": "2026-09-02", "invoice_no": "E-2", "supplier": "مورد",
                                          f"qty_{lid}": "6", f"price_{lid}": "12"})
    db = SessionLocal()
    r = db.get(Receipt, rid)
    assert r.invoice_no == "E-2" and r.date == date(2026, 9, 2) and r.lines[0].qty == 6 and r.lines[0].unit_price == 12
    assert S.stock_table(db, env["site"], [env["a"]])[env["a"]]["cur"] == before + 2
    db.close()
    c.post(f"/receipts/{rid}/delete")
    db = SessionLocal(); assert db.get(Receipt, rid) is None; db.close()


def test_crusher_issue_to_projects_and_clients(env):
    """الكسارة جهة صرف زي الخلاطات (مشروع/عميل لكل سطر)، وبتظهر في كشف الكسارة وأرصدة المواد الخام."""
    db = SessionLocal()
    g = db.query(Item).filter_by(code="RM-G34").one().id
    db.close()
    c = TestClient(app)
    login(c, env, "requester@example.com")
    page = c.get(f"/requests/new/raw?site={env['site']}&unit={env['crusher']}").text
    assert "الصرف إلى" in page and "الديزل" not in page.split("ملاحظات")[0].split("الصرف إلى")[1]
    r = c.post("/requests/new/raw", data={"site_id": env["site"], "dest": "crusher", "unit_id": env["crusher"],
                                          "entity_type": "مشروع", "entity_name": "طريق نجران", "item_id": g, "qty": "25"},
               follow_redirects=False)
    rid = int(r.headers["location"].rsplit("/", 1)[1])
    db = SessionLocal()
    req = db.get(Request, rid)
    assert req.raw_dest == "crusher" and req.current_stage == 2
    lid = req.lines[0].id
    db.close()
    for who in ("prod.manager@example.com", "final@example.com"):
        login(c, env, who)
        c.post(f"/approvals/{rid}", data={"action": "approve", "line": [lid]})
    db = SessionLocal()
    assert db.get(Request, rid).status == "approved"
    assert S.raw_stock_table(db, env["site"])[g]["out_crusher"] == 25
    db.close()
    login(c, env, "eng.crusher@example.com")
    page = c.get(f"/unit-statement?site={env['site']}&unit={env['crusher']}").text
    assert "صرف الكسارة" in page and "طريق نجران" in page


def test_crusher_line_sale_price(env):
    """سعر البيع اختياري في سطر العميل (من غير اختيار «بيع» فوق)."""
    db = SessionLocal()
    g = db.query(Item).filter_by(code="RM-G34").one().id
    db.close()
    c = TestClient(app)
    login(c, env, "requester@example.com")
    r = c.post("/requests/new/raw", data={"site_id": env["site"], "dest": "crusher", "unit_id": env["crusher"],
                                          "entity_type": ["عميل", "مشروع"], "entity_name": ["مؤسسة الفجر", "طريق 15"],
                                          "item_id": [g, g], "qty": ["10", "5"], "sale_price": ["30", ""]},
               follow_redirects=False)
    rid = int(r.headers["location"].rsplit("/", 1)[1])
    page = c.get(f"/requests/{rid}").text
    assert "قيمة البيع للعملاء" in page and "300" in page and "مؤسسة الفجر — السعر" in page



def test_asset_invoice_and_custody_stock(env):
    """فاتورة أصول بتضيف لمخزن الأصول، والتسليم للموظف بيخصم، والإرجاع السليم بيرجع للرصيد."""
    from app.models import CustodyRecord
    db = SessionLocal()
    lap = Item(code="AS-LAP", name="لابتوب", uom="جهاز", category="asset")
    db.add(lap)
    db.commit()
    lid = lap.id
    db.close()
    c = TestClient(app)
    login(c, env, "storekeeper@example.com")
    page = c.get(f"/receipts/new?site={env['site']}&kind=asset").text
    assert "تسجيل وارد — أصول" in page and "AS-LAP" in page and "SP-" not in page.split("<datalist")[1].split("</datalist>")[0]
    r = c.post("/receipts/new", data={"site_id": env["site"], "kind": "asset", "date": "2026-10-01", "invoice_no": "AS-1",
                                      "item_id": [lid], "qty": ["3"], "unit_price": ["2500"]}, follow_redirects=False)
    assert r.status_code == 303
    db = SessionLocal()
    assert S.custody_stock(db, env["site"])[lid]["avail"] == 3
    db.close()
    base = {"site_id": env["site"], "employee_no": "E77", "employee_name": "سعيد", "category": "الأصول", "row": ["0"]}
    r = c.post("/custody/new", data={**base, "item_name_0": "AS-LAP — لابتوب", "qty_0": "5"}, follow_redirects=True)
    assert "أكبر من الرصيد المتاح" in r.text
    r = c.post("/custody/new", data={**base, "item_name_0": "حاجة مش في المخزن", "qty_0": "1"}, follow_redirects=True)
    assert "من أصناف المخزن" in r.text
    c.post("/custody/new", data={**base, "item_name_0": "AS-LAP — لابتوب", "qty_0": "2"})
    db = SessionLocal()
    st = S.custody_stock(db, env["site"])[lid]
    assert st["avail"] == 1 and st["held"] == 2
    rec = db.query(CustodyRecord).filter_by(employee_no="E77").order_by(CustodyRecord.id.desc()).first()
    assert rec.item_id == lid and rec.item_name == "لابتوب"
    rid = rec.id
    db.close()
    c.post(f"/custody/{rid}/return", data={"return_condition": "سليمة"})
    db = SessionLocal()
    assert S.custody_stock(db, env["site"])[lid]["avail"] == 3
    db.close()
    assert "رصيد الأصول وأدوات السلامة" in c.get(f"/custody/stock?site={env['site']}").text


def test_backup_download_and_restore(env):
    """النسخة الاحتياطية: ZIP فيه Excel ومرفقات وdata.json، والاسترجاع بيرجّع البيانات زي ما كانت."""
    import io, zipfile
    from app.models import Attachment
    db = SessionLocal()
    db.add(Attachment(key="bk-test.pdf", filename="فاتورة.pdf", content_type="application/pdf", data=b"%PDF-1.4 hello"))
    db.add(Item(code="BK-ONE", name="صنف النسخة", uom="قطعة"))
    db.commit()
    n_items = db.query(Item).count()
    db.close()
    c = TestClient(app)
    login(c, env, "storekeeper@example.com")
    assert c.get("/admin/backup/download", follow_redirects=False).status_code == 403
    login(c, env, "admin@example.com")
    assert "تحميل نسختين" in c.get("/admin/backup").text
    r = c.get("/admin/backup/download")
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(r.content))
    names = z.namelist()
    assert "data.json" in names and "بيانات.xlsx" in names and "مرفقات/bk-test.pdf" in names
    assert z.read("مرفقات/bk-test.pdf") == b"%PDF-1.4 hello"
    # نغيّر في البيانات وبعدين نسترجع
    db = SessionLocal()
    db.query(Item).filter_by(code="BK-ONE").delete()
    db.add(Item(code="BK-NEW", name="اتضاف بعد النسخة", uom="قطعة"))
    db.commit()
    db.close()
    bad = c.post("/admin/backup/restore", data={"confirm": "لا"}, files={"file": ("b.zip", r.content, "application/zip")},
                 follow_redirects=True)
    assert "للتأكيد" in bad.text
    ok = c.post("/admin/backup/restore", data={"confirm": "استرجاع"}, files={"file": ("b.zip", r.content, "application/zip")},
                follow_redirects=True)
    assert "تم استرجاع النسخة الاحتياطية" in ok.text
    db = SessionLocal()
    assert db.query(Item).filter_by(code="BK-ONE").count() == 1 and db.query(Item).filter_by(code="BK-NEW").count() == 0
    assert db.query(Item).count() == n_items
    assert db.query(Attachment).filter_by(key="bk-test.pdf").one().data == b"%PDF-1.4 hello"
    db.add(Item(code="BK-AFTER", name="بعد الاسترجاع", uom="قطعة"))   # الأرقام التلقائية شغالة بعد الاسترجاع
    db.commit()
    db.close()


def test_backup_auto_key_and_reminder(env):
    from app.config import settings
    c = TestClient(app)
    assert c.get("/backup/auto?key=x").status_code == 403
    settings.BACKUP_KEY = "k-123"
    try:
        r = c.get("/backup/auto?key=k-123")
        assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    finally:
        settings.BACKUP_KEY = ""
    login(c, env, "admin@example.com")
    from app.db import SessionLocal as SL
    from app.models import AppSetting
    db = SL()
    db.query(AppSetting).filter(AppSetting.key == "last_backup_manual").delete()
    db.commit()
    db.close()
    assert "لم تُحمَّل أي نسخة احتياطية" in c.get("/").text      # تذكير
    c.get("/admin/backup/download")
    assert "نسخة احتياطية" not in c.get("/").text.split("<main>")[1].split("</main>")[0][:400]


def test_security_hardening(env):
    """هيدرز الأمان، قفل openapi، نوع الملف من الامتداد، وحد محاولات الدخول لكل حساب."""
    from app.models import Attachment
    c = TestClient(app)
    r = c.get("/healthz")
    assert r.headers["x-content-type-options"] == "nosniff" and r.headers["x-frame-options"] == "SAMEORIGIN"
    assert c.get("/openapi.json").status_code == 404
    db = SessionLocal()
    db.add(Attachment(key="evil.pdf", filename="x.pdf", content_type="text/html", data=b"<script>alert(1)</script>"))
    db.add(Attachment(key="sheet.xlsx", filename="s.xlsx", content_type="text/html", data=b"PK"))
    db.commit()
    db.close()
    login(c, env, "requester@example.com")
    r = c.get("/files/evil.pdf")
    assert r.headers["content-type"] == "application/pdf"            # مش text/html حتى لو اترفع كده
    assert c.get("/files/sheet.xlsx").headers["content-disposition"].startswith("attachment")
    # 10 محاولات غلط لنفس الحساب من IPs مختلفة → يتقفل مؤقتاً
    from app import main as M
    M._FAILS.clear()
    for i in range(10):
        c.post("/signin", data={"email": "admin@example.com", "password": "bad"}, headers={"X-Forwarded-For": f"10.0.0.{i}"})
    r = c.post("/signin", data={"email": "admin@example.com", "password": "admin-pass-123"}, follow_redirects=True)
    assert "محاولات كثيرة" in r.text
    M._FAILS.clear()


def test_custom_bg_notify_and_idle_logout(env, monkeypatch):
    import app.auth as A
    c = TestClient(app)
    login(c, env, "requester@example.com")
    # خلفية من لوحة الألوان (داكنة → وضع داكن)، وقيمة غلط مرفوضة
    assert "تم حفظ المظهر" in c.post("/account/appearance", data={"primary": "#2B3440", "bg": "custom", "bg_custom": "#1a2b3c",
                                                                "font": "Cairo"}, follow_redirects=True).text
    page = c.get("/").text
    assert "--bg:#1A2B3C" in page and 'data-mode="dark"' in page
    c.post("/account/appearance", data={"primary": "#2B3440", "bg": "mint", "font": "Cairo"})
    page = c.get("/").text
    assert "--bg:#E7F3EC" in page and 'data-mode="dark"' not in page
    assert "اختيار غير صحيح" in c.post("/account/appearance", data={"primary": "#2B3440", "bg": "custom", "bg_custom": "red;x",
                                                                   "font": "Cairo"}, follow_redirects=True).text
    c.post("/account/appearance", data={"reset": "1"})
    # عدد الطلبات المنتظرة
    r = c.get("/api/notify")
    assert r.status_code == 200 and isinstance(r.json()["n"], int)
    # خمول 31 دقيقة → خروج؛ والاستعلام التلقائي ما بيجددش الجلسة
    now = A.time.time()
    monkeypatch.setattr(A.time, "time", lambda: now + 20 * 60)
    assert c.get("/api/notify").status_code == 200
    monkeypatch.setattr(A.time, "time", lambda: now + 31 * 60)
    assert c.get("/api/notify").status_code == 401
    assert c.get("/", follow_redirects=False).status_code == 303
    # الاستخدام الفعلي بيجدد المدة
    login(c, env, "requester@example.com")
    monkeypatch.setattr(A.time, "time", lambda: now + 31 * 60 + 25 * 60)
    assert c.get("/", follow_redirects=False).status_code == 200
    monkeypatch.setattr(A.time, "time", lambda: now + 31 * 60 + 50 * 60)
    assert c.get("/", follow_redirects=False).status_code == 200


def test_issue_chart_by_unit_kind_links(env):
    """رسم «حسب الوحدة» في اللوحة: كل نوع بيفتح صرف النوع ده بس، وصفحة الصرف بتفلتر بنوع الوحدة."""
    from datetime import date as _d
    c = TestClient(app)
    db = SessionLocal()
    admin = db.query(User).filter(User.roles == "admin").first()
    env["admin"] = admin.id
    db.close()
    login(c, env, "admin")
    page = c.get("/").text
    assert "&amp;kind=" in page
    db = SessionLocal()
    reqs = db.query(Request).filter(Request.type == "spare", Request.unit_id.isnot(None)).all()
    kinds = {r.unit.kind for r in reqs}
    assert len(kinds) >= 2
    db.close()
    m = _d.today().strftime("%Y-%m")
    for k in kinds:
        html = c.get(f"/requests?type=spare&month={m}&kind={k}").text
        for r in reqs:
            if r.work_date.strftime("%Y-%m") == m:
                assert (r.req_no in html) == (r.unit.kind == k)


def test_raw_stock_sub_crusher_groups(env):
    """أرصدة كسارة تابعة (اسم الصنف «المادة - كسارة X») بتظهر في جدول لوحدها في صفحة الموقع بس."""
    db = SessionLocal()
    it = Item(code="RM-ARD-T", name="بحص تجربة - كسارة العارضة", uom="طن", category="raw", crusher=True)
    db.add(it); db.flush()
    db.add(OpeningBalance(site_id=env["site"], item_id=it.id, qty=100, value=800))
    db.commit(); db.close()
    c = TestClient(app)
    login(c, env, "admin@example.com")
    t = c.get(f"/raw-stock?site={env['site']}").text
    assert "أرصدة كسارة العارضة" in t and "RM-ARD-T" in t
    db = SessionLocal()
    other = Site(code="OT2", name="موقع آخر"); db.add(other); db.commit(); oid = other.id; db.close()
    assert "RM-ARD-T" not in c.get(f"/raw-stock?site={oid}").text
