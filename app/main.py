import io
import json
import uuid
from datetime import date, datetime
from pathlib import Path

from fastapi import Depends, FastAPI, Request
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from sqlalchemy import func, inspect, or_, text
from sqlalchemy.orm import Session
from starlette.middleware.sessions import SessionMiddleware

from . import services as S
from .auth import (Forbidden, LoginRequired, check_password, current_user, hash_password, require,
                   router as auth_router)
from .config import settings
from .db import Base, engine, get_db, SessionLocal
from .models import (CUSTODY_CATEGORIES, CUSTODY_OLD_NAMES, DIESEL_PURPOSES, ENTITY_TYPES, LINE_STATUS, REQ_STATUS, REQ_TYPES, ROLES,
                     STAGES, UNIT_KINDS, ApprovalRoute, Attachment, CustodyRecord, Item, LineDecision, MailLog, OpeningBalance,
                     Receipt, ReceiptLine, Request as Req, RequestLine, Site, Unit, User,
                     TRANSFER_STATUS, Transfer, TransferLine)

APP_DIR = Path(__file__).parent
from contextlib import asynccontextmanager


@asynccontextmanager
async def lifespan(_app):
    init_db()
    yield


app = FastAPI(title=settings.APP_NAME, docs_url=None, redoc_url=None, lifespan=lifespan)
app.add_middleware(SessionMiddleware, secret_key=settings.SECRET_KEY, max_age=60 * 60 * 12,
                   https_only=settings.BASE_URL.startswith("https"), same_site="lax")
app.mount("/static", StaticFiles(directory=APP_DIR / "static"), name="static")
app.include_router(auth_router)
T = Jinja2Templates(directory=APP_DIR / "templates")
T.env.globals.update(fmt=S.fmt, REQ_STATUS=REQ_STATUS, REQ_TYPES=REQ_TYPES, LINE_STATUS=LINE_STATUS, STAGES=STAGES,
                     UNIT_KINDS=UNIT_KINDS, ROLES=ROLES, APP_NAME=settings.APP_NAME, DEV_AUTH=settings.DEV_AUTH,
                     ENTITY_TYPES=ENTITY_TYPES, DIESEL_PURPOSES=DIESEL_PURPOSES, CUSTODY_CATEGORIES=CUSTODY_CATEGORIES,
                     TRANSFER_STATUS=TRANSFER_STATUS)

def _att_url(a):
    return a if a and a.startswith("http") else f"/files/{a}"


T.env.globals["att_url"] = _att_url

ALLOWED_EXT = {".pdf", ".jpg", ".jpeg", ".png", ".webp", ".xlsx", ".xls", ".docx", ".doc"}
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp"}
DIESEL_CODE = "RM-DSL"


# ---------------- تهيئة ----------------
# أعمدة اتضافت لجداول موجودة — create_all ما بيضيفهاش، فبنضيفها هنا مرة واحدة
NEW_COLUMNS = [("custody_records", "photo", "VARCHAR(300)")]


def _add_missing_columns():
    insp = inspect(engine)
    with engine.begin() as conn:
        for table, col, ddl in NEW_COLUMNS:
            if table in insp.get_table_names() and col not in {c["name"] for c in insp.get_columns(table)}:
                conn.execute(text(f"ALTER TABLE {table} ADD COLUMN {col} {ddl}"))


def init_db():
    Base.metadata.create_all(engine)
    _add_missing_columns()
    db = SessionLocal()
    try:
        if not db.query(User).first():
            db.add(User(email=settings.BOOTSTRAP_ADMIN_EMAIL.lower(), name=settings.BOOTSTRAP_ADMIN_NAME,
                        roles="admin", all_sites=True,
                        password_hash=hash_password(settings.BOOTSTRAP_ADMIN_PASSWORD)
                        if settings.BOOTSTRAP_ADMIN_PASSWORD else None))
        raw = [("RM-G34", "بحص 3/4"), ("RM-G38", "بحص 3/8"), ("RM-G1", "بحص 1 بوصة"), ("RM-G15", "بحص 1.5 بوصة"),
               ("RM-PWD", "بودرة"), ("RM-BIT", "بيتومين")]
        for code, name in raw:
            if not db.query(Item).filter(Item.code == code).first():
                db.add(Item(code=code, name=name, uom="طن", category="raw"))
        for old_name, new_name in CUSTODY_OLD_NAMES.items():
            db.query(CustodyRecord).filter(CustodyRecord.category == old_name).update({"category": new_name})
        if not db.query(Item).filter(Item.code == DIESEL_CODE).first():
            db.add(Item(code=DIESEL_CODE, name="ديزل", uom="لتر", category="raw"))
        db.commit()
    finally:
        db.close()


# ---------------- أدوات ----------------
@app.exception_handler(LoginRequired)
def _login_required(request: Request, exc: LoginRequired):
    return RedirectResponse(f"/login?next={request.url.path}", status_code=303)


@app.exception_handler(Forbidden)
def _forbidden(request: Request, exc: Forbidden):
    return HTMLResponse("<meta name=viewport content='width=device-width, initial-scale=1'>"
                        "<div dir=rtl style='font-family:Tahoma;padding:40px'>ليس لديك صلاحية لهذه الصفحة. "
                        "<a href='/'>الرئيسية</a></div>", status_code=403)


def flash(request: Request, msg: str, kind="ok"):
    request.session.setdefault("flash", []).append([kind, msg])


def render(request: Request, name: str, user: User | None = None, **ctx):
    fl = request.session.pop("flash", [])
    pending_n = transfer_n = 0
    if user is not None:
        with SessionLocal() as db:
            u = db.get(User, user.id)
            pending_n = sum(1 for r in db.query(Req).filter(Req.status == "pending").all() if S.can_approve(db, u, r))
            if u.has("storekeeper"):  # تحويلات واردة بانتظار تأكيد الاستلام
                transfer_n = db.query(Transfer).filter(Transfer.status == "in_transit",
                                                       Transfer.to_site_id.in_(u.site_ids(db))).count()
    return T.TemplateResponse(request, name, {"user": user, "flashes": fl, "path": request.url.path,
                                              "pending_n": pending_n, "transfer_n": transfer_n, **ctx})


def back(url):
    return RedirectResponse(url, status_code=303)


def user_sites(db: Session, user: User):
    ids = user.site_ids(db)
    return db.query(Site).filter(Site.id.in_(ids), Site.active.is_(True)).order_by(Site.name).all()


def pick_site(db, user, site_id):
    sites = user_sites(db, user)
    if not sites:
        raise Forbidden()
    try:
        sid = int(site_id) if site_id else None
    except ValueError:
        sid = None
    return sites, next((s for s in sites if s.id == sid), sites[0])


def save_upload(db, upload, images_only=False):
    if upload is None or not getattr(upload, "filename", ""):
        return None
    ext = Path(upload.filename).suffix.lower()
    if images_only and ext not in IMAGE_EXT:
        raise S.BusinessError("الصورة لازم تكون JPG أو PNG")
    if ext not in ALLOWED_EXT:
        raise S.BusinessError("نوع المرفق غير مسموح (PDF أو صورة أو Excel أو Word)")
    data = upload.file.read()
    if len(data) > 10 * 1024 * 1024:
        raise S.BusinessError("حجم المرفق أكبر من 10 ميجا")
    name = f"{uuid.uuid4().hex}{ext}"
    db.add(Attachment(key=name, filename=Path(upload.filename).name[:300],
                      content_type=upload.content_type or "application/octet-stream", data=data))
    return name


def parse_float(v):
    try:
        return float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def parse_date(v, default=None):
    try:
        return date.fromisoformat(v)
    except (TypeError, ValueError):
        return default


def can_view_request(db, user, req):
    if req.site_id in user.site_ids(db) or (req.requester_id == user.id):
        return True
    return any(d.user_id == user.id for d in req.decisions) or S.can_approve(db, user, req)


# ---------------- دخول تجريبي ----------------
@app.get("/dev-login", response_class=HTMLResponse)
def dev_login_page(request: Request, db: Session = Depends(get_db)):
    if not settings.DEV_AUTH:
        return back("/login")
    users = db.query(User).filter(User.active.is_(True)).order_by(User.name).all()
    return render(request, "dev_login.html", users=users)


@app.post("/dev-login")
async def dev_login(request: Request, db: Session = Depends(get_db)):
    if not settings.DEV_AUTH:
        return back("/login")
    form = await request.form()
    u = db.get(User, int(form.get("uid", 0) or 0))
    if not u:
        return back("/dev-login")
    request.session["uid"] = u.id
    return back(request.session.pop("next", "/"))


@app.get("/denied", response_class=HTMLResponse)
def denied(request: Request):
    email = request.session.get("denied_email", "")
    return render(request, "denied.html", email=email)


@app.get("/files/{name}")
def get_file(name: str, db: Session = Depends(get_db), user: User = Depends(current_user)):
    from urllib.parse import quote
    from fastapi.responses import Response
    a = db.query(Attachment).filter_by(key=name).first()
    if not a:
        p = settings.UPLOAD_DIR / Path(name).name
        return FileResponse(p) if p.exists() else HTMLResponse("غير موجود", status_code=404)
    return Response(a.data, media_type=a.content_type,
                    headers={"Content-Disposition": f"inline; filename*=UTF-8''{quote(a.filename)}"})


# ---------------- الدخول بكلمة المرور ----------------
_FAILS: dict = {}


@app.get("/signin", response_class=HTMLResponse)
def signin_page(request: Request):
    return render(request, "signin.html")


@app.post("/signin")
async def signin(request: Request, db: Session = Depends(get_db)):
    import time
    form = await request.form()
    email = (form.get("email") or "").strip().lower()
    ip = request.client.host if request.client else ""
    fails = [t for t in _FAILS.get(ip, []) if time.time() - t < 900]
    if len(fails) >= 10:
        flash(request, "محاولات كثيرة — حاول بعد 15 دقيقة", "err")
        return back("/signin")
    u = db.query(User).filter(User.email == email, User.active.is_(True)).first()
    if not u or not check_password(form.get("password") or "", u.password_hash):
        _FAILS[ip] = fails + [time.time()]
        flash(request, "الإيميل أو كلمة المرور غير صحيحة", "err")
        return back("/signin")
    _FAILS.pop(ip, None)
    nxt = request.session.get("next", "/")
    request.session.clear()
    request.session["uid"] = u.id
    return back(nxt)


@app.get("/account", response_class=HTMLResponse)
def account(request: Request, user: User = Depends(current_user)):
    return render(request, "account.html", user)


@app.post("/account")
async def account_post(request: Request, db: Session = Depends(get_db), user: User = Depends(current_user)):
    form = await request.form()
    u = db.get(User, user.id)
    new = form.get("new") or ""
    if u.password_hash and not check_password(form.get("old") or "", u.password_hash):
        flash(request, "كلمة المرور الحالية غير صحيحة", "err")
    elif len(new) < 8 or new != form.get("new2"):
        flash(request, "كلمة المرور الجديدة يجب ألا تقل عن 8 أحرف وأن تتطابق مع التأكيد", "err")
    else:
        u.password_hash = hash_password(new)
        db.commit()
        flash(request, "تم تغيير كلمة المرور")
    return back("/account")


# ---------------- لوحة المتابعة ----------------
@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request, site: str | None = None, month: str | None = None,
              db: Session = Depends(get_db), user: User = Depends(current_user)):
    sites = user_sites(db, user)
    site_ids = [s.id for s in sites]
    sel = next((s for s in sites if str(s.id) == str(site)), None)
    ids = [sel.id] if sel else site_ids
    today = date.today()
    try:
        y, m = (int(x) for x in (month or f"{today:%Y-%m}").split("-"))
    except ValueError:
        y, m = today.year, today.month
    start = date(y, m, 1)
    end = date(y + (m == 12), m % 12 + 1, 1)

    my_pending = [r for r in db.query(Req).filter(Req.status == "pending").all() if S.can_approve(db, user, r)]
    base = db.query(Req).filter(Req.site_id.in_(ids))
    pending_by_stage = {s: base.filter(Req.status == "pending", Req.current_stage == s).count() for s in STAGES}
    month_reqs = base.filter(Req.work_date >= start, Req.work_date < end).count()

    # قيمة المصروف المعتمد (قطع غيار) حسب الوحدة
    spend_rows = (db.query(Site.name, Unit.name, func.sum(RequestLine.value), func.count(func.distinct(Req.id)))
                  .select_from(RequestLine).join(Req).join(Site, Site.id == Req.site_id)
                  .outerjoin(Unit, Unit.id == Req.unit_id)
                  .filter(Req.site_id.in_(ids), Req.type == "spare", Req.status.in_(["approved", "partial"]),
                          RequestLine.status == "approved", Req.work_date >= start, Req.work_date < end)
                  .group_by(Site.name, Unit.name).order_by(func.sum(RequestLine.value).desc()).all())
    month_spend = sum(r[2] or 0 for r in spend_rows)

    top_items = (db.query(Item.code, Item.name, Item.uom, func.sum(RequestLine.qty), func.sum(RequestLine.value))
                 .select_from(RequestLine).join(Req).join(Item)
                 .filter(Req.site_id.in_(ids), Req.type == "spare", Req.status.in_(["approved", "partial"]),
                         RequestLine.status == "approved", Req.work_date >= start, Req.work_date < end)
                 .group_by(Item.code, Item.name, Item.uom).order_by(func.sum(RequestLine.value).desc()).limit(10).all())

    raw_rows = (db.query(Item.name, Item.uom, RequestLine.diesel_purpose, func.sum(RequestLine.qty))
                .select_from(RequestLine).join(Req).join(Item)
                .filter(Req.site_id.in_(ids), Req.type == "raw", Req.status.in_(["approved", "partial"]),
                        RequestLine.status == "approved", Req.work_date >= start, Req.work_date < end)
                .group_by(Item.name, Item.uom, RequestLine.diesel_purpose).order_by(Item.name).all())

    stock_value = 0.0
    for sid in ids:
        stock_value += sum(r["value"] for r in S.stock_table(db, sid).values())

    open_custody = db.query(CustodyRecord).filter(CustodyRecord.site_id.in_(ids),
                                                  CustodyRecord.returned_at.is_(None)).count()
    recent = base.order_by(Req.created_at.desc()).limit(8).all()
    return render(request, "dashboard.html", user, sites=sites, sel=sel, month=f"{y:04d}-{m:02d}",
                  my_pending=my_pending, pending_by_stage=pending_by_stage, month_reqs=month_reqs,
                  month_spend=month_spend, spend_rows=spend_rows, top_items=top_items, raw_rows=raw_rows,
                  stock_value=stock_value, open_custody=open_custody, recent=recent)


# ---------------- API ----------------
@app.get("/api/stock")
def api_stock(site: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    if site not in user.site_ids(db):
        raise Forbidden()
    st = S.stock_table(db, site)
    items = db.query(Item).filter(Item.category == "spare", Item.active.is_(True)).order_by(Item.code).all()
    iss = S.issued_summary(db, site, [i.id for i in items], date.today())
    return JSONResponse([{"id": i.id, "code": i.code, "name": i.name, "uom": i.uom,
                          "available": round(st[i.id]["available"], 3) if i.id in st else 0,
                          "cur": round(st[i.id]["cur"], 3) if i.id in st else 0,
                          "m_qty": round(iss[i.id]["m_qty"], 3) if i.id in iss else 0,
                          "m_n": iss[i.id]["m_n"] if i.id in iss else 0,
                          "y_qty": round(iss[i.id]["y_qty"], 3) if i.id in iss else 0,
                          "y_n": iss[i.id]["y_n"] if i.id in iss else 0} for i in items])


# ---------------- الطلبات ----------------
@app.get("/requests", response_class=HTMLResponse)
def requests_list(request: Request, type: str = "", status: str = "", site: str = "", q: str = "",
                  mine: str = "", page: int = 1, db: Session = Depends(get_db), user: User = Depends(current_user)):
    sites = user_sites(db, user)
    qry = db.query(Req).filter(or_(Req.site_id.in_([s.id for s in sites]), Req.requester_id == user.id))
    if type in REQ_TYPES:
        qry = qry.filter(Req.type == type)
    if status in REQ_STATUS:
        qry = qry.filter(Req.status == status)
    if site.isdigit():
        qry = qry.filter(Req.site_id == int(site))
    if mine:
        qry = qry.filter(Req.requester_id == user.id)
    if q.strip():
        like = f"%{q.strip()}%"
        qry = qry.filter(or_(Req.req_no.ilike(like), Req.reason.ilike(like), Req.requester_name.ilike(like),
                             Req.lines.any(RequestLine.item.has(or_(Item.name.ilike(like), Item.code.ilike(like))))))
    total = qry.count()
    per = 50
    rows = qry.order_by(Req.work_date.desc(), Req.id.desc()).offset((page - 1) * per).limit(per).all()
    return render(request, "requests_list.html", user, rows=rows, total=total, page=page, pages=(total + per - 1) // per,
                  sites=sites, f=dict(type=type, status=status, site=site, q=q, mine=mine))


@app.get("/requests/new/spare", response_class=HTMLResponse)
def new_spare(request: Request, site: str | None = None, db: Session = Depends(get_db),
              user: User = Depends(require("requester", "storekeeper"))):
    sites, sel = pick_site(db, user, site)
    units = [u for u in sel.units if u.active]
    return render(request, "request_spare.html", user, sites=sites, sel=sel, units=units, today=date.today())


@app.post("/requests/new/spare")
async def new_spare_post(request: Request, db: Session = Depends(get_db),
                         user: User = Depends(require("requester", "storekeeper"))):
    form = await request.form()
    sites, sel = pick_site(db, user, form.get("site_id"))
    try:
        unit = db.get(Unit, int(form.get("unit_id") or 0))
        if unit and unit.site_id != sel.id:
            unit = None
        lines = []
        for iid, qty in zip(form.getlist("item_id"), form.getlist("qty")):
            if not iid and not qty:
                continue
            item = db.get(Item, int(iid)) if str(iid).isdigit() else None
            if not item or item.category != "spare":
                raise S.BusinessError("اختر الصنف من القائمة")
            lines.append(dict(item=item, qty=parse_float(qty)))
        req = S.create_request(db, rtype="spare", site=sel, unit=unit, requester=user,
                               work_date=parse_date(form.get("work_date"), date.today()),
                               reason=form.get("reason"), lines=lines)
    except S.BusinessError as e:
        db.rollback()
        flash(request, str(e), "err")
        return back(f"/requests/new/spare?site={sel.id}")
    flash(request, f"تم إرسال الطلب {req.req_no} للاعتماد")
    return back(f"/requests/{req.id}")


@app.get("/requests/new/raw", response_class=HTMLResponse)
def new_raw(request: Request, site: str | None = None, db: Session = Depends(get_db),
            user: User = Depends(require("requester", "storekeeper"))):
    sites, sel = pick_site(db, user, site)
    units = [u for u in sel.units if u.active and u.kind == "asphalt"]
    items = db.query(Item).filter(Item.category == "raw", Item.code != DIESEL_CODE, Item.active.is_(True)).order_by(Item.id).all()
    return render(request, "request_raw.html", user, sites=sites, sel=sel, units=units, items=items, today=date.today())


@app.post("/requests/new/raw")
async def new_raw_post(request: Request, db: Session = Depends(get_db),
                       user: User = Depends(require("requester", "storekeeper"))):
    form = await request.form()
    sites, sel = pick_site(db, user, form.get("site_id"))
    try:
        unit = db.get(Unit, int(form.get("unit_id") or 0))
        if not unit or unit.site_id != sel.id or unit.kind != "asphalt":
            raise S.BusinessError("اختر خلاطة الاسفلت")
        lines = []
        for et, en, iid, qty in zip(form.getlist("entity_type"), form.getlist("entity_name"),
                                    form.getlist("item_id"), form.getlist("qty")):
            if not (en or "").strip() and not qty:
                continue
            item = db.get(Item, int(iid)) if str(iid).isdigit() else None
            if not item or item.category != "raw" or item.code == DIESEL_CODE:
                raise S.BusinessError("اختر المادة من القائمة")
            if et not in ENTITY_TYPES:
                raise S.BusinessError("اختر نوع الجهة (مشروع / عميل)")
            if not (en or "").strip():
                raise S.BusinessError("اكتب اسم الجهة")
            lines.append(dict(item=item, qty=parse_float(qty), entity_type=et, entity_name=en.strip()))
        diesel = db.query(Item).filter(Item.code == DIESEL_CODE).one()
        for purpose in DIESEL_PURPOSES:
            q = parse_float(form.get(f"diesel_{purpose}"))
            if q:
                lines.append(dict(item=diesel, qty=q, diesel_purpose=purpose))
        req = S.create_request(db, rtype="raw", site=sel, unit=unit, requester=user,
                               work_date=parse_date(form.get("work_date"), date.today()),
                               reason=form.get("reason"), lines=lines)
    except S.BusinessError as e:
        db.rollback()
        flash(request, str(e), "err")
        return back(f"/requests/new/raw?site={sel.id}")
    flash(request, f"تم إرسال الطلب {req.req_no} للاعتماد")
    return back(f"/requests/{req.id}")


@app.get("/requests/{rid}", response_class=HTMLResponse)
def request_detail(rid: int, request: Request, db: Session = Depends(get_db), user: User = Depends(current_user)):
    req = db.get(Req, rid)
    if not req or not can_view_request(db, user, req):
        raise Forbidden()
    approvers = S.approvers_for(db, req, req.current_stage) if req.status == "pending" else []
    st, iss = _spare_ctx(db, req)
    return render(request, "request_detail.html", user, r=req, approvers=approvers, st=st, iss=iss,
                  can_approve=S.can_approve(db, user, req),
                  can_cancel=req.status == "pending" and req.requester_id == user.id and not req.decisions)


@app.post("/requests/{rid}/cancel")
def request_cancel(rid: int, request: Request, db: Session = Depends(get_db), user: User = Depends(current_user)):
    req = db.get(Req, rid)
    if req and req.status == "pending" and req.requester_id == user.id and not req.decisions:
        req.status, req.closed_at = "cancelled", datetime.now()
        db.commit()
        flash(request, "تم إلغاء الطلب")
    return back(f"/requests/{rid}")


# ---------------- الاعتماد ----------------
def _spare_ctx(db, req):
    """الرصيد + المنصرف خلال شهر وسنة الطلب لكل صنف (لطلبات قطع الغيار)."""
    if req.type != "spare":
        return {}, {}
    ids = [l.item_id for l in req.lines]
    return (S.stock_table(db, req.site_id, ids),
            S.issued_summary(db, req.site_id, ids, req.work_date or date.today(), exclude_request_id=req.id))


@app.get("/approvals", response_class=HTMLResponse)
def approvals(request: Request, db: Session = Depends(get_db), user: User = Depends(current_user)):
    rows = [r for r in db.query(Req).filter(Req.status == "pending").order_by(Req.created_at).all()
            if S.can_approve(db, user, r)]
    done = (db.query(Req).join(LineDecision).filter(LineDecision.user_id == user.id)
            .distinct().order_by(Req.id.desc()).limit(30).all())
    return render(request, "approvals.html", user, rows=rows, done=done)


@app.get("/approvals/{rid}", response_class=HTMLResponse)
def approval_page(rid: int, request: Request, db: Session = Depends(get_db), user: User = Depends(current_user)):
    req = db.get(Req, rid)
    if not req:
        raise Forbidden()
    if not S.can_approve(db, user, req):
        if can_view_request(db, user, req):
            flash(request, "هذا الطلب ليس بانتظار اعتمادك الآن (تم البت فيه أو في مرحلة أخرى)", "err")
            return back(f"/requests/{rid}")
        raise Forbidden()
    st, iss = _spare_ctx(db, req)
    return render(request, "approval.html", user, r=req, st=st, iss=iss)


def _apply_decision(request, db, req, user, form):
    """يرجع (True, رسالة) عند النجاح أو (False, رسالة الخطأ)."""
    action = form.get("action")
    ids = set() if action == "reject_all" else {int(x) for x in form.getlist("line") if str(x).isdigit()}
    comment = form.get("comment")
    if action != "reject_all" and not ids:
        return False, "لم تحدد أي صنف للموافقة — لرفض الطلب كاملاً استخدم زر «رفض الكل»"
    has_rejects = any(l.status == "pending" and l.id not in ids for l in req.lines)
    if has_rejects and not (comment or "").strip():
        return False, "اكتب سبب الرفض في الملاحظات"
    try:
        S.decide(db, req, user, ids, comment)
    except S.BusinessError as e:
        db.rollback()
        return False, str(e)
    return True, f"تم تسجيل قرارك على الطلب {req.req_no}"


@app.post("/approvals/{rid}")
async def approval_post(rid: int, request: Request, db: Session = Depends(get_db), user: User = Depends(current_user)):
    req = db.get(Req, rid)
    if not req:
        raise Forbidden()
    ok, msg = _apply_decision(request, db, req, user, await request.form())
    flash(request, msg, "ok" if ok else "err")
    return back("/approvals" if ok else f"/approvals/{rid}")


# ---------------- الاعتماد من رابط الإيميل (بدون تسجيل دخول) ----------------
def _token_ctx(db, token):
    """يرجع (req, approver, رسالة خطأ)."""
    data = S.read_approval_token(token)
    if not data:
        return None, None, "الرابط غير صالح أو انتهت صلاحيته — افتح النظام وادخل من «الاعتمادات»"
    rid, stage, uid = data
    req, approver = db.get(Req, rid), db.get(User, uid)
    if not req or not approver or not approver.active:
        return None, None, "الرابط غير صالح"
    if req.status != "pending" or req.current_stage != stage or not S.can_approve(db, approver, req):
        return req, approver, f"تم البت في الطلب {req.req_no} في هذه المرحلة — الحالة الحالية: {REQ_STATUS[req.status]}"
    return req, approver, None


def _token_page(request, db, req, approver, token, do="", msg=None, kind="ok"):
    if msg:
        return render(request, "approval_done.html", None, r=req, approver=approver, msg=msg, kind=kind)
    st, iss = _spare_ctx(db, req)
    return render(request, "approval.html", None, r=req, st=st, iss=iss, approver=approver, token=token, do=do)


@app.get("/a/{token}", response_class=HTMLResponse)
def approval_link(token: str, request: Request, do: str = "", db: Session = Depends(get_db)):
    # GET لا يغيّر شيئاً (برامج فحص الروابط في Outlook تفتح الروابط تلقائياً) — القرار بزر في الصفحة
    req, approver, err = _token_ctx(db, token)
    if err:
        return _token_page(request, db, req, approver, token, msg=err, kind="err")
    return _token_page(request, db, req, approver, token, do=do)


@app.post("/a/{token}", response_class=HTMLResponse)
async def approval_link_post(token: str, request: Request, db: Session = Depends(get_db)):
    req, approver, err = _token_ctx(db, token)
    if err:
        return _token_page(request, db, req, approver, token, msg=err, kind="err")
    form = await request.form()
    ok, msg = _apply_decision(request, db, req, approver, form)
    if not ok:
        flash(request, msg, "err")
        return _token_page(request, db, req, approver, token, do=form.get("action", ""))
    db.refresh(req)
    nxt = (f" — انتقل إلى مرحلة {STAGES[req.current_stage]}" if req.status == "pending"
           else f" — الحالة: {REQ_STATUS[req.status]}")
    return _token_page(request, db, req, approver, token, msg=msg + nxt)


# ---------------- الأرصدة ----------------
@app.get("/stock", response_class=HTMLResponse)
def stock(request: Request, site: str | None = None, q: str = "", only: str = "", db: Session = Depends(get_db),
          user: User = Depends(current_user)):
    sites, sel = pick_site(db, user, site)
    st = S.stock_table(db, sel.id)
    items = {i.id: i for i in db.query(Item).filter(Item.id.in_(list(st.keys()))).all()}
    rows = []
    for iid, r in st.items():
        it = items.get(iid)
        if not it or it.category != "spare":
            continue
        if q and q.strip().lower() not in it.name.lower() and q.strip().lower() not in it.code.lower():
            continue
        if only == "zero" and r["cur"] > 0:
            continue
        if only == "has" and r["cur"] <= 0:
            continue
        rows.append((it, r))
    rows.sort(key=lambda x: x[0].code)
    total_value = sum(r["value"] for _, r in rows)
    return render(request, "stock.html", user, sites=sites, sel=sel, rows=rows, q=q, only=only, total_value=total_value)


@app.get("/stock/{item_id}", response_class=HTMLResponse)
def item_card(item_id: int, request: Request, site: str | None = None, db: Session = Depends(get_db),
              user: User = Depends(current_user)):
    sites, sel = pick_site(db, user, site)
    item = db.get(Item, item_id)
    if not item:
        raise Forbidden()
    moves = []
    ob = db.query(OpeningBalance).filter_by(site_id=sel.id, item_id=item_id).first()
    if ob:
        moves.append(dict(d=None, kind="رصيد افتتاحي", ref="", inq=ob.qty, outq=0, price=(ob.value / ob.qty) if ob.qty else 0, who=""))
    for rl in (db.query(ReceiptLine).join(Receipt).filter(Receipt.site_id == sel.id, ReceiptLine.item_id == item_id)):
        moves.append(dict(d=rl.receipt.date, kind="وارد", ref=rl.receipt.invoice_no or "", inq=rl.qty, outq=0,
                          price=rl.unit_price, who=rl.receipt.supplier or "", link=f"/receipts/{rl.receipt_id}",
                          att=rl.receipt.attachment, value=(rl.qty or 0) * (rl.unit_price or 0)))
    for l in (db.query(RequestLine).join(Req).filter(Req.site_id == sel.id, RequestLine.item_id == item_id,
                                                      RequestLine.status == "approved",
                                                      Req.status.in_(["approved", "partial"]))):
        moves.append(dict(d=l.request.work_date, kind="صرف", ref=l.request.req_no, inq=0, outq=l.qty, price=l.unit_cost,
                          who=(l.request.unit.name if l.request.unit else ""), link=f"/requests/{l.request_id}",
                          value=l.value))
    for tl in (db.query(TransferLine).join(Transfer).filter(
            TransferLine.item_id == item_id, Transfer.status.in_(["in_transit", "received"]),
            or_(Transfer.from_site_id == sel.id, Transfer.to_site_id == sel.id))):
        t = tl.transfer
        if t.from_site_id == sel.id:
            q = tl.qty if t.status == "in_transit" else (tl.recv_qty or 0)
            moves.append(dict(d=t.date, kind="تحويل صادر", ref=t.tr_no, inq=0, outq=q, price=tl.unit_cost,
                              who=f"إلى {t.to_site.name}" + (" (في الطريق)" if t.status == "in_transit" else ""),
                              link=f"/transfers/{t.id}", value=round(q * (tl.unit_cost or 0), 2)))
        elif t.status == "received":
            q = tl.recv_qty or 0
            moves.append(dict(d=t.date, kind="تحويل وارد", ref=t.tr_no, inq=q, outq=0, price=tl.unit_cost,
                              who=f"من {t.from_site.name}", link=f"/transfers/{t.id}", value=round(q * (tl.unit_cost or 0), 2)))
    moves.sort(key=lambda m: (m["d"] is not None, m["d"] or date.min))
    bal = 0
    for m in moves:
        bal += m["inq"] - m["outq"]
        m["bal"] = bal
    st = S.stock_table(db, sel.id, [item_id]).get(item_id)
    return render(request, "item_card.html", user, sites=sites, sel=sel, item=item, moves=moves, st=st)


# ---------------- الوارد ----------------
@app.get("/receipts", response_class=HTMLResponse)
def receipts(request: Request, site: str | None = None, q: str = "", db: Session = Depends(get_db),
             user: User = Depends(current_user)):
    sites, sel = pick_site(db, user, site)
    qry = db.query(Receipt).filter(Receipt.site_id == sel.id)
    if q.strip():
        like = f"%{q.strip()}%"
        qry = qry.filter(or_(Receipt.invoice_no.ilike(like), Receipt.supplier.ilike(like),
                             Receipt.lines.any(ReceiptLine.item.has(or_(Item.name.ilike(like), Item.code.ilike(like))))))
    rows = qry.order_by(Receipt.date.desc(), Receipt.id.desc()).limit(500).all()
    return render(request, "receipts.html", user, sites=sites, sel=sel, rows=rows, q=q)


@app.post("/receipts/{rid}/attach")
async def receipt_attach(rid: int, request: Request, db: Session = Depends(get_db),
                         user: User = Depends(require("storekeeper"))):
    rec = db.get(Receipt, rid)
    if not rec or rec.site_id not in user.site_ids(db):
        raise Forbidden()
    form = await request.form()
    try:
        key = save_upload(db, form.get("file"))
    except S.BusinessError as e:
        flash(request, str(e), "err")
        return back(f"/receipts/{rid}")
    if not key:
        flash(request, "اختر ملف الفاتورة", "err")
        return back(f"/receipts/{rid}")
    rec.attachment = key
    db.commit()
    flash(request, "تم رفع الفاتورة")
    return back(f"/receipts/{rid}")


@app.get("/receipts/new", response_class=HTMLResponse)
def receipt_new(request: Request, site: str | None = None, db: Session = Depends(get_db),
                user: User = Depends(require("storekeeper"))):
    sites, sel = pick_site(db, user, site)
    items = db.query(Item).filter(Item.category == "spare", Item.active.is_(True)).order_by(Item.code).all()
    return render(request, "receipt_new.html", user, sites=sites, sel=sel, items=items, today=date.today())


@app.post("/receipts/new")
async def receipt_new_post(request: Request, db: Session = Depends(get_db), user: User = Depends(require("storekeeper"))):
    form = await request.form()
    sites, sel = pick_site(db, user, form.get("site_id"))
    try:
        rec = Receipt(site_id=sel.id, date=parse_date(form.get("date"), date.today()),
                      invoice_no=(form.get("invoice_no") or "").strip() or None,
                      supplier=(form.get("supplier") or "").strip() or None, created_by_id=user.id)
        for iid, qty, price in zip(form.getlist("item_id"), form.getlist("qty"), form.getlist("unit_price")):
            if not iid and not qty:
                continue
            item = db.get(Item, int(iid)) if str(iid).isdigit() else None
            q, p = parse_float(qty), parse_float(price)
            if not item:
                raise S.BusinessError("اختر الصنف من القائمة")
            if not q or q <= 0 or p is None or p < 0:
                raise S.BusinessError(f"الكمية/السعر غير صحيح للصنف «{item.name}»")
            rec.lines.append(ReceiptLine(item_id=item.id, qty=q, unit_price=p))
        if not rec.lines:
            raise S.BusinessError("أضف صنفاً واحداً على الأقل")
        rec.attachment = save_upload(db, form.get("attachment"))
        db.add(rec)
        db.commit()
    except S.BusinessError as e:
        db.rollback()
        flash(request, str(e), "err")
        return back(f"/receipts/new?site={sel.id}")
    flash(request, "تم تسجيل الوارد")
    return back(f"/receipts/{rec.id}")


@app.get("/receipts/{rid}", response_class=HTMLResponse)
def receipt_detail(rid: int, request: Request, db: Session = Depends(get_db), user: User = Depends(current_user)):
    rec = db.get(Receipt, rid)
    if not rec or rec.site_id not in user.site_ids(db):
        raise Forbidden()
    return render(request, "receipt_detail.html", user, r=rec)


# ---------------- التحويل بين المواقع ----------------
@app.get("/transfers", response_class=HTMLResponse)
def transfers(request: Request, site: str | None = None, db: Session = Depends(get_db),
              user: User = Depends(require("storekeeper", "transfer"))):
    sites, sel = pick_site(db, user, site)
    rows = (db.query(Transfer).filter(or_(Transfer.from_site_id == sel.id, Transfer.to_site_id == sel.id))
            .order_by((Transfer.status == "in_transit").desc(), Transfer.date.desc(), Transfer.id.desc()).limit(500).all())
    return render(request, "transfers.html", user, sites=sites, sel=sel, rows=rows)


# إنشاء التحويل: صلاحية «تحويل بين الفروع» + الربط بالموقع المرسِل فقط (مش لازم الربط بالموقع المستلِم).
# تأكيد الاستلام: أمين مستودع مربوط بالموقع المستلِم.
@app.get("/transfers/new", response_class=HTMLResponse)
def transfer_new(request: Request, site: str | None = None, db: Session = Depends(get_db),
                 user: User = Depends(require("transfer"))):
    sites, sel = pick_site(db, user, site)
    targets = db.query(Site).filter(Site.active.is_(True), Site.id != sel.id).order_by(Site.name).all()
    return render(request, "transfer_new.html", user, sites=sites, sel=sel, targets=targets, today=date.today())


@app.post("/transfers/new")
async def transfer_new_post(request: Request, db: Session = Depends(get_db), user: User = Depends(require("transfer"))):
    form = await request.form()
    sites, sel = pick_site(db, user, form.get("from_site_id"))
    try:
        to_site = db.get(Site, int(form.get("to_site_id") or 0))
        if to_site and not to_site.active:
            to_site = None
        lines = []
        for iid, qty in zip(form.getlist("item_id"), form.getlist("qty")):
            if not iid and not qty:
                continue
            item = db.get(Item, int(iid)) if str(iid).isdigit() else None
            if not item or item.category != "spare":
                raise S.BusinessError("اختر الصنف من القائمة")
            lines.append(dict(item=item, qty=parse_float(qty)))
        tr = S.create_transfer(db, from_site=sel, to_site=to_site, user=user,
                               tr_date=parse_date(form.get("date"), date.today()), note=form.get("note"), lines=lines)
    except S.BusinessError as e:
        db.rollback()
        flash(request, str(e), "err")
        return back(f"/transfers/new?site={sel.id}")
    flash(request, f"تم إرسال التحويل {tr.tr_no} — بانتظار تأكيد الاستلام من {tr.to_site.name}")
    return back(f"/transfers/{tr.id}")


def _transfer_for(db, user, tid):
    tr = db.get(Transfer, tid)
    ids = user.site_ids(db)
    if not tr or (tr.from_site_id not in ids and tr.to_site_id not in ids):
        raise Forbidden()
    return tr, ids


@app.get("/transfers/{tid}", response_class=HTMLResponse)
def transfer_detail(tid: int, request: Request, db: Session = Depends(get_db), user: User = Depends(require("storekeeper", "transfer"))):
    tr, ids = _transfer_for(db, user, tid)
    return render(request, "transfer_detail.html", user, t=tr,
                  can_receive=tr.status == "in_transit" and tr.to_site_id in ids and user.has("storekeeper"),
                  can_cancel=tr.status == "in_transit" and tr.from_site_id in ids and user.has("transfer"))


@app.post("/transfers/{tid}/receive")
async def transfer_receive(tid: int, request: Request, db: Session = Depends(get_db),
                           user: User = Depends(require("storekeeper"))):
    tr, ids = _transfer_for(db, user, tid)
    if tr.to_site_id not in ids:
        raise Forbidden()
    form = await request.form()
    try:
        S.receive_transfer(db, tr, user, {l.id: parse_float(form.get(f"recv_{l.id}")) for l in tr.lines}, form.get("note"))
    except S.BusinessError as e:
        db.rollback()
        flash(request, str(e), "err")
        return back(f"/transfers/{tid}")
    flash(request, f"تم تأكيد استلام التحويل {tr.tr_no} وإضافته لرصيد {tr.to_site.name}")
    return back(f"/transfers/{tid}")


@app.post("/transfers/{tid}/cancel")
def transfer_cancel(tid: int, request: Request, db: Session = Depends(get_db), user: User = Depends(require("transfer"))):
    tr, ids = _transfer_for(db, user, tid)
    if tr.from_site_id not in ids:
        raise Forbidden()
    try:
        S.cancel_transfer(db, tr)
    except S.BusinessError as e:
        flash(request, str(e), "err")
        return back(f"/transfers/{tid}")
    flash(request, f"تم إلغاء التحويل {tr.tr_no} ورجعت الكميات لرصيد {tr.from_site.name}")
    return back(f"/transfers/{tid}")


# ---------------- العهد ----------------
@app.get("/custody", response_class=HTMLResponse)
def custody(request: Request, site: str | None = None, q: str = "", state: str = "open", cat: str = "",
            db: Session = Depends(get_db), user: User = Depends(current_user)):
    sites, sel = pick_site(db, user, site)
    qry = db.query(CustodyRecord).filter(CustodyRecord.site_id == sel.id)
    if cat not in CUSTODY_CATEGORIES:
        cat = ""
    counts = dict(db.query(CustodyRecord.category, func.count()).filter(
        CustodyRecord.site_id == sel.id, CustodyRecord.returned_at.is_(None)).group_by(CustodyRecord.category).all())
    if cat:
        qry = qry.filter(CustodyRecord.category == cat)
    if state == "open":
        qry = qry.filter(CustodyRecord.returned_at.is_(None))
    elif state == "returned":
        qry = qry.filter(CustodyRecord.returned_at.is_not(None))
    if q.strip():
        like = f"%{q.strip()}%"
        qry = qry.filter(or_(CustodyRecord.employee_no.like(like), CustodyRecord.employee_name.like(like),
                             CustodyRecord.item_name.like(like), CustodyRecord.serial_no.like(like)))
    rows = qry.order_by(CustodyRecord.issued_at.desc(), CustodyRecord.id.desc()).limit(500).all()
    return render(request, "custody.html", user, sites=sites, sel=sel, rows=rows, q=q, state=state, cat=cat,
                  counts=counts)


@app.get("/custody/new", response_class=HTMLResponse)
def custody_new(request: Request, site: str | None = None, emp: str = "", cat: str = "", db: Session = Depends(get_db),
                user: User = Depends(require("storekeeper"))):
    sites, sel = pick_site(db, user, site)
    name = ""
    if emp:
        last = db.query(CustodyRecord).filter_by(employee_no=emp).order_by(CustodyRecord.id.desc()).first()
        name = last.employee_name if last else ""
    return render(request, "custody_new.html", user, sites=sites, sel=sel, today=date.today(), emp=emp, emp_name=name,
                  cat=cat)


@app.post("/custody/new")
async def custody_new_post(request: Request, db: Session = Depends(get_db), user: User = Depends(require("storekeeper"))):
    form = await request.form()
    sites, sel = pick_site(db, user, form.get("site_id"))
    try:
        emp_no = (form.get("employee_no") or "").strip()
        emp_name = (form.get("employee_name") or "").strip()
        cat = form.get("category")
        if not emp_no or not emp_name:
            raise S.BusinessError("الرقم الوظيفي واسم الموظف إجباريان")
        if cat not in CUSTODY_CATEGORIES:
            raise S.BusinessError("اختر النوع (الأصول / أدوات السلامة)")
        att = save_upload(db, form.get("attachment"))
        n = 0
        for k in form.getlist("row"):  # كل سطر ليه رقم، وصورته اسمها photo_<رقم>
            name = (form.get(f"item_name_{k}") or "").strip()
            if not name:
                continue
            q = parse_float(form.get(f"qty_{k}")) or 1
            photo = save_upload(db, form.get(f"photo_{k}"), images_only=True)
            db.add(CustodyRecord(site_id=sel.id, employee_no=emp_no, employee_name=emp_name, category=cat,
                                 item_name=name, qty=q, serial_no=(form.get(f"serial_no_{k}") or "").strip() or None,
                                 issued_at=parse_date(form.get("issued_at"), date.today()), issued_by_id=user.id,
                                 attachment=att, photo=photo, notes=(form.get("notes") or "").strip() or None))
            n += 1
        if not n:
            raise S.BusinessError("أضف بنداً واحداً على الأقل")
        db.commit()
    except S.BusinessError as e:
        db.rollback()
        flash(request, str(e), "err")
        return back(f"/custody/new?site={sel.id}")
    flash(request, f"تم تسجيل {n} بند")
    return back(f"/custody/employee/{emp_no}?site={sel.id}")


@app.get("/custody/employee/{emp_no}", response_class=HTMLResponse)
def custody_employee(emp_no: str, request: Request, site: str | None = None, db: Session = Depends(get_db),
                     user: User = Depends(current_user)):
    sites, sel = pick_site(db, user, site)
    rows = (db.query(CustodyRecord).filter(CustodyRecord.employee_no == emp_no,
                                           CustodyRecord.site_id.in_([s.id for s in sites]))
            .order_by(CustodyRecord.returned_at.is_not(None), CustodyRecord.issued_at.desc()).all())
    return render(request, "custody_employee.html", user, sites=sites, sel=sel, rows=rows, emp_no=emp_no,
                  emp_name=rows[0].employee_name if rows else "", today=date.today())


@app.post("/custody/{cid}/photo")
async def custody_photo(cid: int, request: Request, db: Session = Depends(get_db),
                        user: User = Depends(require("storekeeper"))):
    rec = db.get(CustodyRecord, cid)
    if not rec or rec.site_id not in user.site_ids(db):
        raise Forbidden()
    form = await request.form()
    try:
        key = save_upload(db, form.get("photo"), images_only=True)
    except S.BusinessError as e:
        flash(request, str(e), "err")
        return back(f"/custody/employee/{rec.employee_no}?site={rec.site_id}")
    if key:
        rec.photo = key
        db.commit()
        flash(request, f"تم رفع صورة «{rec.item_name}»")
    return back(f"/custody/employee/{rec.employee_no}?site={rec.site_id}")


@app.post("/custody/{cid}/return")
async def custody_return(cid: int, request: Request, db: Session = Depends(get_db),
                         user: User = Depends(require("storekeeper"))):
    rec = db.get(CustodyRecord, cid)
    if not rec or rec.site_id not in user.site_ids(db):
        raise Forbidden()
    form = await request.form()
    if rec.returned_at is None:
        rec.returned_at = parse_date(form.get("returned_at"), date.today())
        rec.return_condition = (form.get("return_condition") or "").strip() or None
        rec.return_notes = (form.get("return_notes") or "").strip() or None
        db.commit()
        flash(request, f"تم تسجيل إرجاع «{rec.item_name}»")
    return back(f"/custody/employee/{rec.employee_no}?site={rec.site_id}")


# ---------------- الإدارة ----------------
@app.get("/admin/users", response_class=HTMLResponse)
def admin_users(request: Request, edit: int | None = None, db: Session = Depends(get_db), user: User = Depends(require("admin"))):
    users = db.query(User).order_by(User.active.desc(), User.name).all()
    return render(request, "admin_users.html", user, users=users, sites=db.query(Site).order_by(Site.name).all(),
                  e=db.get(User, edit) if edit else None)


@app.post("/admin/users")
async def admin_users_post(request: Request, db: Session = Depends(get_db), user: User = Depends(require("admin"))):
    form = await request.form()
    uid = form.get("id")
    u = db.get(User, int(uid)) if uid else User()
    email = (form.get("email") or "").strip().lower()
    if not email or not (form.get("name") or "").strip():
        flash(request, "الإيميل والاسم إجباريان", "err")
        return back("/admin/users")
    dup = db.query(User).filter(User.email == email, User.id != (u.id or 0)).first()
    if dup:
        flash(request, "الإيميل مسجل لمستخدم آخر", "err")
        return back("/admin/users")
    u.email, u.name = email, form.get("name").strip()
    u.employee_no = (form.get("employee_no") or "").strip() or None
    u.roles = ",".join(r for r in form.getlist("roles") if r in ROLES) or "requester"
    u.all_sites = bool(form.get("all_sites"))
    u.active = bool(form.get("active"))
    u.sites = db.query(Site).filter(Site.id.in_([int(x) for x in form.getlist("sites") if x.isdigit()])).all()
    pw = form.get("password") or ""
    if pw:
        if len(pw) < 8:
            flash(request, "كلمة المرور يجب ألا تقل عن 8 أحرف", "err")
            return back("/admin/users")
        u.password_hash = hash_password(pw)
    db.add(u)
    db.commit()
    flash(request, "تم الحفظ")
    return back("/admin/users")


@app.get("/admin/sites", response_class=HTMLResponse)
def admin_sites(request: Request, db: Session = Depends(get_db), user: User = Depends(require("admin"))):
    return render(request, "admin_sites.html", user, sites=db.query(Site).order_by(Site.name).all())


@app.post("/admin/sites")
async def admin_sites_post(request: Request, db: Session = Depends(get_db), user: User = Depends(require("admin"))):
    form = await request.form()
    act = form.get("act")
    if act == "site":
        code, name = (form.get("code") or "").strip().upper(), (form.get("name") or "").strip()
        if code and name and not db.query(Site).filter_by(code=code).first():
            db.add(Site(code=code, name=name))
        else:
            flash(request, "الكود والاسم إجباريان والكود لا يتكرر", "err")
    elif act == "unit":
        sid, name, kind = form.get("site_id"), (form.get("name") or "").strip(), form.get("kind")
        if sid and name and kind in UNIT_KINDS and not db.query(Unit).filter_by(site_id=int(sid), name=name).first():
            db.add(Unit(site_id=int(sid), name=name, kind=kind))
        else:
            flash(request, "بيانات الوحدة غير مكتملة أو مكررة", "err")
    elif act == "toggle_unit":
        u = db.get(Unit, int(form.get("id")))
        if u:
            u.active = not u.active
    elif act == "toggle_site":
        s = db.get(Site, int(form.get("id")))
        if s:
            s.active = not s.active
    db.commit()
    return back("/admin/sites")


@app.get("/admin/routes", response_class=HTMLResponse)
def admin_routes(request: Request, db: Session = Depends(get_db), user: User = Depends(require("admin"))):
    routes = db.query(ApprovalRoute).order_by(ApprovalRoute.stage, ApprovalRoute.unit_kind).all()
    users = db.query(User).filter(User.active.is_(True)).order_by(User.name).all()
    return render(request, "admin_routes.html", user, routes=routes, users=users,
                  sites=db.query(Site).order_by(Site.name).all())


@app.post("/admin/routes")
async def admin_routes_post(request: Request, db: Session = Depends(get_db), user: User = Depends(require("admin"))):
    form = await request.form()
    if form.get("act") == "delete":
        r = db.get(ApprovalRoute, int(form.get("id")))
        if r:
            db.delete(r)
    else:
        stage = int(form.get("stage") or 0)
        uid = int(form.get("user_id") or 0)
        kind = form.get("unit_kind") or None
        sid = int(form.get("site_id")) if (form.get("site_id") or "").isdigit() else None
        if stage not in STAGES or not db.get(User, uid) or (kind and kind not in UNIT_KINDS):
            flash(request, "بيانات غير مكتملة", "err")
            return back("/admin/routes")
        db.add(ApprovalRoute(stage=stage, user_id=uid, unit_kind=kind if stage == 1 else None, site_id=sid))
    db.commit()
    flash(request, "تم الحفظ")
    return back("/admin/routes")


@app.get("/admin/items", response_class=HTMLResponse)
def admin_items(request: Request, q: str = "", edit: int | None = None, db: Session = Depends(get_db),
                user: User = Depends(require("admin", "storekeeper"))):
    qry = db.query(Item)
    if q:
        qry = qry.filter(or_(Item.code.like(f"%{q}%"), Item.name.like(f"%{q}%")))
    return render(request, "admin_items.html", user, items=qry.order_by(Item.category.desc(), Item.code).limit(500).all(),
                  q=q, e=db.get(Item, edit) if edit else None)


@app.post("/admin/items")
async def admin_items_post(request: Request, db: Session = Depends(get_db),
                           user: User = Depends(require("admin", "storekeeper"))):
    form = await request.form()
    iid = form.get("id")
    it = db.get(Item, int(iid)) if iid else Item()
    code, name = (form.get("code") or "").strip(), (form.get("name") or "").strip()
    if not code or not name:
        flash(request, "الكود والاسم إجباريان", "err")
        return back("/admin/items")
    if db.query(Item).filter(Item.code == code, Item.id != (it.id or 0)).first():
        flash(request, "الكود موجود مسبقاً", "err")
        return back("/admin/items")
    if it.id and it.code == DIESEL_CODE:
        code = DIESEL_CODE
    it.code, it.name = code, name
    it.uom = (form.get("uom") or "قطعة").strip()
    it.category = form.get("category") if form.get("category") in ("spare", "raw") else "spare"
    it.active = bool(form.get("active"))
    db.add(it)
    db.commit()
    flash(request, "تم الحفظ")
    return back("/admin/items")


@app.get("/admin/opening", response_class=HTMLResponse)
def admin_opening(request: Request, site: str | None = None, db: Session = Depends(get_db),
                  user: User = Depends(require("admin"))):
    sites, sel = pick_site(db, user, site)
    rows = (db.query(OpeningBalance).join(Item).filter(OpeningBalance.site_id == sel.id).order_by(Item.code).all())
    items = db.query(Item).filter(Item.category == "spare").order_by(Item.code).all()
    return render(request, "admin_opening.html", user, sites=sites, sel=sel, rows=rows, items=items)


@app.post("/admin/opening")
async def admin_opening_post(request: Request, db: Session = Depends(get_db), user: User = Depends(require("admin"))):
    form = await request.form()
    sites, sel = pick_site(db, user, form.get("site_id"))
    iid = int(form.get("item_id") or 0)
    qty, val = parse_float(form.get("qty")), parse_float(form.get("value"))
    if not db.get(Item, iid) or qty is None or val is None:
        flash(request, "بيانات غير مكتملة", "err")
        return back(f"/admin/opening?site={sel.id}")
    ob = db.query(OpeningBalance).filter_by(site_id=sel.id, item_id=iid).first() or OpeningBalance(site_id=sel.id, item_id=iid)
    ob.qty, ob.value = qty, val
    ob.location = (form.get("location") or "").strip() or ob.location
    db.add(ob)
    db.commit()
    flash(request, "تم الحفظ")
    return back(f"/admin/opening?site={sel.id}")


@app.get("/admin/import", response_class=HTMLResponse)
def admin_import(request: Request, user: User = Depends(require("admin"))):
    return render(request, "admin_import.html", user)


@app.post("/admin/import")
async def admin_import_post(request: Request, user: User = Depends(require("admin"))):
    from scripts.import_excel import run as run_import
    form = await request.form()
    f = form.get("file")
    code, name = (form.get("site_code") or "").strip().upper(), (form.get("site_name") or "").strip()
    if not f or not getattr(f, "filename", "") or not code or not name:
        flash(request, "اختر الملف واكتب كود الموقع واسمه", "err")
        return back("/admin/import")
    try:
        msg = run_import(io.BytesIO(f.file.read()), code, name)
    except Exception as e:  # noqa: BLE001
        flash(request, f"فشل الاستيراد: {e}", "err")
        return back("/admin/import")
    flash(request, msg)
    return back("/admin/import")


# ---------------- رفع الفواتير بالجملة ----------------
def _inv_norm(x):
    import re as _re
    return _re.sub(r"[^0-9a-z\u0600-\u06ff]", "", str(x or "").lower()).lstrip("0")


def _inv_tokens(name):
    import re as _re
    stem = Path(name).stem
    return {_inv_norm(t) for t in _re.split(r"[\s_\-–—]+", stem) if _inv_norm(t)} | {_inv_norm(stem)}


@app.get("/admin/invoices", response_class=HTMLResponse)
def admin_invoices(request: Request, db: Session = Depends(get_db), user: User = Depends(require("storekeeper"))):
    sites = user_sites(db, user)
    stats = []
    for s in sites:
        total = db.query(Receipt).filter(Receipt.site_id == s.id).count()
        inside = db.query(Receipt).filter(Receipt.site_id == s.id, Receipt.attachment.isnot(None),
                                          ~Receipt.attachment.like("http%")).count()
        stats.append((s, total, inside))
    return render(request, "admin_invoices.html", user, sites=sites, stats=stats, result=None)


@app.post("/admin/invoices", response_class=HTMLResponse)
async def admin_invoices_post(request: Request, db: Session = Depends(get_db),
                              user: User = Depends(require("storekeeper"))):
    form = await request.form()
    sites = user_sites(db, user)
    try:
        sid = int(form.get("site") or 0)
    except ValueError:
        sid = 0
    site = next((s for s in sites if s.id == sid), None)
    if not site:
        raise Forbidden()
    recs = [r for r in db.query(Receipt).filter(Receipt.site_id == site.id).all() if r.invoice_no]
    by_inv = {}
    for r in recs:
        by_inv.setdefault(_inv_norm(r.invoice_no), []).append(r)
    matched, unmatched, errors = [], [], []
    for f in form.getlist("files"):
        if not getattr(f, "filename", ""):
            continue
        toks = _inv_tokens(f.filename)
        hits = [k for k in by_inv if k and k in toks]
        if not hits:  # أرقام طويلة موجودة جوه اسم الملف
            flat = _inv_norm(Path(f.filename).stem)
            hits = [k for k in by_inv if len(k) >= 6 and k in flat]
        if not hits:
            unmatched.append(f.filename)
            continue
        key = max(hits, key=len)
        try:
            att = save_upload(db, f)
        except S.BusinessError as e:
            errors.append(f"{f.filename}: {e}")
            continue
        for r in by_inv[key]:
            r.attachment = att
        matched.append((f.filename, by_inv[key][0].invoice_no, len(by_inv[key])))
    db.commit()
    stats = []
    for s in sites:
        total = db.query(Receipt).filter(Receipt.site_id == s.id).count()
        inside = db.query(Receipt).filter(Receipt.site_id == s.id, Receipt.attachment.isnot(None),
                                          ~Receipt.attachment.like("http%")).count()
        stats.append((s, total, inside))
    return render(request, "admin_invoices.html", user, sites=sites, stats=stats,
                  result=dict(site=site, matched=matched, unmatched=unmatched, errors=errors))


@app.get("/admin/mail", response_class=HTMLResponse)
def admin_mail(request: Request, show: int | None = None, db: Session = Depends(get_db),
               user: User = Depends(require("admin"))):
    rows = db.query(MailLog).order_by(MailLog.id.desc()).limit(100).all()
    return render(request, "admin_mail.html", user, rows=rows, cur=db.get(MailLog, show) if show else None)


# ---------------- تصدير Excel ----------------
def _xlsx(headers, rows, title):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B2A5B")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(list(r))
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(str(h)) + 6)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _send_xlsx(buf, fname):
    from urllib.parse import quote
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


@app.get("/export/requests.xlsx")
def export_requests(type: str = "spare", site: str = "", frm: str = "", to: str = "",
                    db: Session = Depends(get_db), user: User = Depends(current_user)):
    ids = user.site_ids(db)
    qry = (db.query(RequestLine).join(Req).filter(Req.site_id.in_(ids), Req.type == (type if type in REQ_TYPES else "spare")))
    if site.isdigit():
        qry = qry.filter(Req.site_id == int(site))
    if parse_date(frm):
        qry = qry.filter(Req.work_date >= parse_date(frm))
    if parse_date(to):
        qry = qry.filter(Req.work_date <= parse_date(to))
    rows = []
    for l in qry.order_by(Req.work_date, Req.id, RequestLine.id):
        r = l.request
        base = [r.req_no, r.work_date, r.site.name, r.unit.name if r.unit else "", l.item.code, l.item.name, l.qty, l.item.uom]
        if r.type == "raw":
            base += [l.entity_type or "", l.entity_name or "", l.diesel_purpose or ""]
        else:
            base += [l.unit_cost, l.value]
        base += [r.requester_display, r.reason or "", REQ_STATUS[r.status], LINE_STATUS[l.status],
                 STAGES.get(l.rejected_stage, "") if l.rejected_stage else ""]
        rows.append(base)
    h = ["رقم الطلب", "التاريخ", "الموقع", "الوحدة", "كود الصنف", "اسم الصنف", "الكمية", "الوحدة"]
    h += ["نوع الجهة", "اسم الجهة", "غرض الديزل"] if type == "raw" else ["سعر الوحدة", "القيمة"]
    h += ["طالب الصرف", "السبب", "حالة الطلب", "حالة الصنف", "رُفض في مرحلة"]
    return _send_xlsx(_xlsx(h, rows, "الصرف"), f"صرف_{REQ_TYPES.get(type, '')}.xlsx")


@app.get("/export/stock.xlsx")
def export_stock(site: str | None = None, db: Session = Depends(get_db), user: User = Depends(current_user)):
    sites, sel = pick_site(db, user, site)
    st = S.stock_table(db, sel.id)
    items = {i.id: i for i in db.query(Item).filter(Item.id.in_(list(st.keys())), Item.category == "spare")}
    rows = [[it.code, it.name, it.uom, r["location"] or "", r["open_qty"], r["in_qty"], r["out_qty"], r["cur"],
             round(r["avg"], 4), round(r["value"], 2), r["pending_qty"], r["tr_in_qty"], r["tr_out_qty"]]
            for iid, r in sorted(st.items(), key=lambda x: items[x[0]].code if x[0] in items else "")
            if (it := items.get(iid))]
    h = ["كود الصنف", "اسم الصنف", "الوحدة", "الموقع", "الرصيد الافتتاحي", "إجمالي الوارد", "المنصرف المعتمد",
         "الرصيد الحالي", "متوسط السعر", "قيمة الرصيد", "قيد الاعتماد", "تحويل وارد", "تحويل صادر"]
    return _send_xlsx(_xlsx(h, rows, "الرصيد"), f"رصيد_{sel.name}.xlsx")


@app.get("/export/custody.xlsx")
def export_custody(site: str | None = None, db: Session = Depends(get_db), user: User = Depends(current_user)):
    sites, sel = pick_site(db, user, site)
    rows = [[c.employee_no, c.employee_name, c.category, c.item_name, c.qty, c.serial_no or "", c.issued_at,
             c.issued_by.name if c.issued_by else "", c.returned_at or "", c.return_condition or "", c.notes or ""]
            for c in db.query(CustodyRecord).filter_by(site_id=sel.id).order_by(CustodyRecord.employee_no, CustodyRecord.issued_at)]
    h = ["الرقم الوظيفي", "اسم الموظف", "النوع", "البند", "الكمية", "الرقم التسلسلي", "تاريخ التسليم", "سلّمها",
         "تاريخ الإرجاع", "الحالة عند الإرجاع", "ملاحظات"]
    return _send_xlsx(_xlsx(h, rows, "الأصول"), f"أصول_{sel.name}.xlsx")


# ---------------- روابط Power BI (CSV) ----------------
import csv as _csv
import hmac as _hmac


def _csv_response(head, rows, name):
    buf = io.StringIO()
    buf.write("\ufeff")
    w = _csv.writer(buf)
    w.writerow(head)
    w.writerows(rows)
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv; charset=utf-8",
                             headers={"Content-Disposition": f'inline; filename="{name}.csv"'})


@app.get("/export/{table}.csv")
def export_csv(table: str, key: str = "", db: Session = Depends(get_db)):
    if not settings.EXPORT_KEY or not _hmac.compare_digest(key, settings.EXPORT_KEY):
        return JSONResponse({"error": "forbidden"}, status_code=403)
    st_req = {"pending": "قيد الاعتماد", "approved": "معتمد نهائي", "partial": "معتمد جزئياً",
              "rejected": "مرفوض", "cancelled": "ملغي"}
    st_line = {"pending": "قيد الاعتماد", "approved": "معتمد نهائي", "rejected": "مرفوض"}
    if table == "issues":
        q = (db.query(RequestLine, Req).join(Req, RequestLine.request_id == Req.id)
             .order_by(Req.work_date, Req.id, RequestLine.id))
        rows = []
        for l, r in q:
            rows.append([r.site.name, r.req_no, "قطع غيار" if r.type == "spare" else "مواد خام",
                         r.work_date.isoformat() if r.work_date else "", r.unit.name if r.unit else "",
                         l.item.code, l.item.name, l.item.uom or "", l.qty, l.entity_type or "", l.entity_name or "",
                         l.diesel_purpose or "", r.requester_display, r.reason or "",
                         st_line.get(l.status, l.status), st_req.get(r.status, r.status), r.current_stage,
                         l.unit_cost if l.unit_cost is not None else "", l.value if l.value is not None else ""])
        head = ["المستودع", "رقم الطلب", "نوع الطلب", "التاريخ", "الوحدة", "كود الصنف", "اسم الصنف", "وحدة القياس",
                "الكمية", "نوع الجهة", "الجهة", "بند الديزل", "طالب الصرف", "السبب", "حالة الصنف", "حالة الطلب",
                "المرحلة", "سعر الوحدة", "القيمة"]
        return _csv_response(head, rows, "issues")
    if table == "stock":
        rows = []
        items = {i.id: i for i in db.query(Item)}
        for site in db.query(Site).order_by(Site.id):
            last = dict(db.query(RequestLine.item_id, func.max(Req.work_date)).join(Req)
                        .filter(Req.site_id == site.id, RequestLine.status == "approved", Req.type == "spare")
                        .group_by(RequestLine.item_id).all())
            for iid, t in S.stock_table(db, site.id).items():
                it = items.get(iid)
                if not it:
                    continue
                rows.append([site.name, it.code, it.name, it.uom or "", round(t["open_qty"], 4), round(t["in_qty"], 4),
                             round(t["out_qty"], 4), round(t["pending_qty"], 4), round(t["cur"], 4),
                             round(t["avg"], 4), round(t["value"], 2),
                             last[iid].isoformat() if last.get(iid) else ""])
        head = ["المستودع", "كود الصنف", "اسم الصنف", "وحدة القياس", "الرصيد الافتتاحي", "إجمالي الوارد",
                "المنصرف المعتمد", "تحت الاعتماد", "الرصيد الحالي", "متوسط السعر", "قيمة الرصيد", "آخر تاريخ صرف"]
        return _csv_response(head, rows, "stock")
    if table == "receipts":
        q = db.query(ReceiptLine, Receipt).join(Receipt).order_by(Receipt.date, Receipt.id)
        rows = [[r.site.name, r.date.isoformat() if r.date else "", r.invoice_no or "", r.supplier or "",
                 l.item.code, l.item.name, l.qty, l.unit_price, round((l.qty or 0) * (l.unit_price or 0), 2)]
                for l, r in q]
        head = ["المستودع", "التاريخ", "رقم الفاتورة", "المورد", "كود الصنف", "اسم الصنف", "الكمية", "سعر الوحدة",
                "الإجمالي"]
        return _csv_response(head, rows, "receipts")
    return JSONResponse({"error": "unknown table"}, status_code=404)
