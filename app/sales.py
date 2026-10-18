"""المبيعات (مستخلصات العملاء الشهرية): الاستيراد من Excel وتجميع لوحة المبيعات.
المبيعات هنا مستقلة عن أرصدة المخزون — مصدرها إيميلات المستخلصات المعتمدة."""
import io
from collections import defaultdict
from datetime import date, datetime

from sqlalchemy.orm import Session

from .models import SALES_KINDS, SalesLine, Site

PAL = ["#4F6582", "#0f9d58", "#f5b301", "#e8590c", "#7c3aed", "#0891b2", "#db2777", "#65a30d", "#9ca3af"]
KIND_COLOR = {"crusher": PAL[0], "asphalt": PAL[2], "concrete": PAL[5], "other": PAL[8]}
MONTHS_AR = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]
HEAD = ["الشهر", "النوع", "الوحدة", "العميل", "الصنف", "وحدة القياس", "الكمية", "السعر", "المبلغ بدون ضريبة",
        "الضريبة", "الإجمالي", "ملاحظات", "المصدر"]
KIND_BY_AR = {v: k for k, v in SALES_KINDS.items()} | {"كسارة": "crusher", "خلاطة أسفلت": "asphalt", "أسفلت": "asphalt",
                                                        "اسفلت": "asphalt", "خرسانه": "concrete"}
# الوحدة ← الموقع (بالاسم)
SITE_HINTS = [("الضيق", "نجران"), ("نجران", "نجران"), ("الكدم", "الكدم"), ("مك", "مك"), ("الجنادري", "الجنادري"),
              ("الدار", "الدار"), ("المدين", "الدار")]


def _norm(s):
    return " ".join(str(s or "").replace("ـ", "").split())


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("ر.س.", "").replace("‏", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def _period(v):
    if isinstance(v, datetime):
        return date(v.year, v.month, 1)
    if isinstance(v, date):
        return v.replace(day=1)
    s = str(v or "").strip()
    for sep in ("-", "/"):
        p = s.split(sep)
        if len(p) >= 2 and p[0].isdigit() and p[1].isdigit():
            y, m = (int(p[0]), int(p[1])) if len(p[0]) == 4 else (int(p[1]), int(p[0]))
            return date(y, m, 1)
    raise ValueError(f"الشهر غير مفهوم: {v!r} (اكتبه 2026-09)")


def guess_site(db: Session, unit: str):
    n = _norm(unit).replace("ى", "ي").replace("ة", "ه")
    sites = db.query(Site).filter(Site.active.is_(True)).all()
    for hint, name in SITE_HINTS:
        if hint.replace("ة", "ه") in n:
            for s in sites:
                if name.replace("ة", "ه") in s.name.replace("ى", "ي").replace("ة", "ه"):
                    return s.id
    return None


def import_xlsx(db: Session, data: bytes, source_default=""):
    """يستورد شيت «المبيعات» (أو أول شيت). كل (شهر، وحدة) في الملف بيستبدل المسجّل قبل كده لنفس الشهر والوحدة."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb["المبيعات"] if "المبيعات" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("الملف فاضي")
    head = [_norm(h) for h in rows[0]]
    missing = [h for h in ("الشهر", "الوحدة", "العميل", "المبلغ بدون ضريبة") if h not in head]
    if missing:
        raise ValueError("أعمدة ناقصة: " + "، ".join(missing))
    ix = {h: i for i, h in enumerate(head)}
    get = lambda r, h: r[ix[h]] if h in ix and ix[h] < len(r) else None
    lines, keys = [], set()
    for n, r in enumerate(rows[1:], start=2):
        if not any(v not in (None, "") for v in r):
            continue
        unit, cust = _norm(get(r, "الوحدة")), _norm(get(r, "العميل"))
        amount = _num(get(r, "المبلغ بدون ضريبة"))
        if not unit or not cust or amount is None:
            raise ValueError(f"سطر {n}: الوحدة والعميل والمبلغ مطلوبين")
        period = _period(get(r, "الشهر"))
        kind = KIND_BY_AR.get(_norm(get(r, "النوع")), _norm(get(r, "النوع")) if _norm(get(r, "النوع")) in SALES_KINDS else None)
        if not kind:
            u = unit.replace("ة", "ه")
            kind = "crusher" if "كسار" in u else "asphalt" if ("اسفلت" in u or "أسفلت" in u) else "concrete" if "خرسان" in u else "other"
        vat = _num(get(r, "الضريبة"))
        total = _num(get(r, "الإجمالي"))
        if vat is None:
            vat = round((total - amount), 2) if total is not None else 0.0
        if total is None:
            total = round(amount + vat, 2)
        qty = _num(get(r, "الكمية")) or 0.0
        price = _num(get(r, "السعر"))
        if price is None and qty:
            price = round(amount / qty, 4)
        lines.append(SalesLine(period=period, kind=kind, unit=unit, site_id=guess_site(db, unit), customer=cust,
                               product=_norm(get(r, "الصنف")) or None, uom=_norm(get(r, "وحدة القياس")) or None,
                               qty=qty, price=price, amount=amount, vat=vat, total=total,
                               notes=_norm(get(r, "ملاحظات"))[:200] or None,
                               source=(_norm(get(r, "المصدر")) or source_default)[:300] or None))
        keys.add((period, unit))
    for period, unit in keys:
        db.query(SalesLine).filter(SalesLine.period == period, SalesLine.unit == unit).delete()
    db.add_all(lines)
    return len(lines), len(keys)


def months_list(f: date, t: date):
    out, y, m = [], f.year, f.month
    while (y, m) <= (t.year, t.month):
        out.append((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def _hbar(cid, title, rows, color, unit="", n=10, span=1):
    rows = sorted(rows, key=lambda r: -r[1])[:n]
    return dict(id=cid, type="bar", horizontal=True, title=title, unit=unit, span=span, labels=[r[0] for r in rows],
                datasets=[dict(label=title, data=[round(r[1], 2) for r in rows], backgroundColor=color, borderRadius=6)])


def _donut(cid, title, rows, unit="", n=7):
    rows = sorted([r for r in rows if r[1] > 0], key=lambda r: -r[1])
    if len(rows) > n:
        rows = rows[:n - 1] + [("أخرى", sum(r[1] for r in rows[n - 1:]))]
    return dict(id=cid, type="doughnut", title=title, unit=unit, labels=[r[0] for r in rows],
                datasets=[dict(data=[round(r[1], 2) for r in rows], backgroundColor=PAL[:len(rows)])])


def build(lines, f: date, t: date):
    """lines: SalesLine مفلترة → أرقام اللوحة والرسوم والجداول."""
    months = months_list(f, t)
    k = dict(amount=0.0, vat=0.0, total=0.0, tons=0.0, m3=0.0)
    custs, units = set(), set()
    per_m_kind = defaultdict(lambda: defaultdict(float))
    by_unit = defaultdict(lambda: dict(kind="", amount=0.0, total=0.0, qty=defaultdict(float), custs=set()))
    by_cust = defaultdict(lambda: dict(amount=0.0, total=0.0, qty=defaultdict(float), units=set()))
    by_prod = defaultdict(lambda: dict(qty=0.0, amount=0.0, uom=""))
    unit_month = defaultdict(lambda: defaultdict(float))
    for ln in lines:
        ym = (ln.period.year, ln.period.month)
        uom = ln.uom or ""
        k["amount"] += ln.amount or 0
        k["vat"] += ln.vat or 0
        k["total"] += ln.total or 0
        if "م" in uom and "3" in uom or uom in ("م3", "م٣", "متر مكعب"):
            k["m3"] += ln.qty or 0
        elif uom.startswith("طن") or uom == "":
            k["tons"] += ln.qty or 0
        custs.add(ln.customer)
        units.add(ln.unit)
        per_m_kind[ym][ln.kind] += ln.amount or 0
        u = by_unit[ln.unit]
        u["kind"] = ln.kind
        u["amount"] += ln.amount or 0
        u["total"] += ln.total or 0
        u["qty"][uom or "طن"] += ln.qty or 0
        u["custs"].add(ln.customer)
        c = by_cust[ln.customer]
        c["amount"] += ln.amount or 0
        c["total"] += ln.total or 0
        c["qty"][uom or "طن"] += ln.qty or 0
        c["units"].add(ln.unit)
        if ln.product:
            p = by_prod[(ln.product, uom)]
            p["qty"] += ln.qty or 0
            p["amount"] += ln.amount or 0
        unit_month[ln.unit][ym] += ln.amount or 0
    kinds = [kk for kk in SALES_KINDS if any(per_m_kind[ym].get(kk) for ym in months)]
    labels = [f"{MONTHS_AR[m - 1]} {y}" if (len(months) > 12 or m == 1) else MONTHS_AR[m - 1] for y, m in months]
    charts = [
        dict(id="s1", type="bar", stacked=True, title="المبيعات الشهرية حسب النوع (ريال بدون ضريبة)", unit="ريال", span=2,
             labels=labels, datasets=[dict(label=SALES_KINDS[kk], data=[round(per_m_kind[ym].get(kk, 0), 2) for ym in months],
                                           backgroundColor=KIND_COLOR[kk]) for kk in kinds]),
        _donut("s2", "المبيعات حسب النوع", [(SALES_KINDS[kk], sum(per_m_kind[ym].get(kk, 0) for ym in months)) for kk in kinds], "ريال"),
        _hbar("s3", "المبيعات حسب الوحدة", [(u, v["amount"]) for u, v in by_unit.items()], PAL[1], "ريال"),
        _hbar("s4", "أكبر 10 عملاء", [(cc, v["amount"]) for cc, v in by_cust.items()], PAL[4], "ريال", span=2),
    ]
    unit_rows = sorted(by_unit.items(), key=lambda x: -x[1]["amount"])
    cust_rows = sorted(by_cust.items(), key=lambda x: -x[1]["amount"])
    prod_rows = sorted(by_prod.items(), key=lambda x: -x[1]["amount"])
    pivot = [(u, [unit_month[u].get(ym, 0) for ym in months], sum(unit_month[u].values())) for u, _ in unit_rows]
    return dict(k=k, n_cust=len(custs), n_units=len(units), months=months, month_labels=labels, charts=charts,
                units=unit_rows, custs=cust_rows, prods=prod_rows, pivot=pivot,
                month_tot=[sum(per_m_kind[ym].values()) for ym in months])
