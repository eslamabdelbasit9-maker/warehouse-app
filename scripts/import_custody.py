"""استيراد سجلات العهد (الأصول / أدوات السلامة) من شيتات Excel الحالية.

بيقرا أي شيت فيه عمود «الرقم الوظيفي» + («اسم الأصل» أو «المهمات المستلمة»):
  - «اسم الأصل» ← نوع «الأصول»، «المهمات المستلمة» ← نوع «أدوات السلامة».
  - الفرع بيتربط بالموقع بالاسم (الكدمي ← الكدمى جيزان، الضيقة ← نجران، ...).
  - روابط الصورة والمرفق (OneDrive) بتتحفظ زي ما هي.
  - السطر المكرر (نفس الموقع + الرقم الوظيفي + البند + النوع، ولسه ما رجعش) بيتخطى، فإعادة الاستيراد آمنة.
"""
import re
from datetime import date

from openpyxl import load_workbook

from app.db import SessionLocal
from app.models import CustodyRecord, Site
from scripts.import_excel import clean, num, to_date

BRANCH_ALIASES = {"الضيقه": "نجران", "الضيقة": "نجران"}


def norm(s):
    s = clean(s) or ""
    s = s.replace("ة", "ه").replace("ى", "ي").replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    return re.sub(r"\s+", " ", s).strip()


def find_site(sites, branch):
    b = norm(BRANCH_ALIASES.get(clean(branch) or "", branch))
    if not b:
        return None
    for s in sites:
        n = norm(s.name)
        if n == b or b in n.split(" ") or n.startswith(b) or b.startswith(n):
            return s
    return None


def _date(v):
    if isinstance(v, str):
        v = v.strip().replace("/", "-")
        try:
            return to_date(v)
        except ValueError:
            return None
    try:
        return to_date(v)
    except (TypeError, ValueError):
        return None


def _link(cell):
    if cell is None:
        return None
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target[:300]
    v = clean(cell.value)
    return v[:300] if v and v.startswith("http") else None


def run(fileobj, user_id=None):
    wb = load_workbook(fileobj)
    db = SessionLocal()
    sites = db.query(Site).filter(Site.active.is_(True)).all()
    added = skipped = 0
    missing_sites = set()
    no_date = 0
    try:
        for ws in wb.worksheets:
            head = [clean(c.value) for c in ws[1]]
            if "الرقم الوظيفي" not in head:
                continue
            if "اسم الأصل" in head:
                cat, item_col = "الأصول", head.index("اسم الأصل")
            elif "المهمات المستلمة" in head:
                cat, item_col = "أدوات السلامة", head.index("المهمات المستلمة")
            else:
                continue
            col = {h: i for i, h in enumerate(head) if h}
            for row in ws.iter_rows(min_row=2):
                def v(name):
                    i = col.get(name)
                    return clean(row[i].value) if i is not None and row[i].value is not None else None
                emp_no, item = v("الرقم الوظيفي"), clean(row[item_col].value) if row[item_col].value else None
                if not emp_no or not item:
                    continue
                emp_no = emp_no[:-2] if emp_no.endswith(".0") else emp_no
                site = find_site(sites, v("الفرع"))
                if not site:
                    missing_sites.add(v("الفرع") or "—")
                    continue
                dup = (db.query(CustodyRecord).filter_by(site_id=site.id, employee_no=emp_no, item_name=item[:250],
                                                         category=cat).filter(CustodyRecord.returned_at.is_(None)).first())
                if dup:
                    skipped += 1
                    continue
                d = _date(row[col["تاريخ الاستلام"]].value) if "تاريخ الاستلام" in col else None
                notes = [x for x in (v("الوظيفة"), v("القسم")) if x]
                if not d:
                    no_date += 1
                    notes.append("تاريخ الاستلام مش مسجل في الشيت")
                code = v("كود الأصل")
                db.add(CustodyRecord(
                    site_id=site.id, employee_no=emp_no, employee_name=v("اسم الموظف") or emp_no, category=cat,
                    item_name=item[:250], qty=num(v("الكمية")) or 1,
                    serial_no=code if code and code not in ("-", "—") else None,
                    issued_at=d or date.today(), issued_by_id=user_id,
                    attachment=_link(row[col["رابط المرفق"]]) if "رابط المرفق" in col else None,
                    photo=_link(row[col["رابط الصورة"]]) if "رابط الصورة" in col else None,
                    notes=" — ".join(notes) or None))
                added += 1
        db.commit()
    finally:
        db.close()
    msg = f"تم: اتسجل {added} بند عهدة"
    if skipped:
        msg += f"، واتخطّى {skipped} متسجل قبل كده"
    if no_date:
        msg += f" — {no_date} بند من غير تاريخ استلام اتسجل بتاريخ النهارده"
    if missing_sites:
        msg += f" — فروع مش لاقي لها موقع: {'، '.join(sorted(missing_sites))}"
    return msg
