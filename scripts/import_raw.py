"""استيراد حركة المواد الخام لموقع من ملف Excel موحّد (بيتجهز من تقارير الإنتاج وتقارير الخلاطة).

شيتات الملف:
  الأصناف:  الكود | الاسم | الوحدة | ناتج كسارة (1/0)
  الافتتاحي: الكود | الكمية | القيمة | التاريخ (اختياري)
  الإنتاج:  التاريخ | الكود | الكمية | السعر
  الوارد:   التاريخ | رقم الفاتورة | المورد | الكود | الكمية | السعر | بدون ضريبة (1/0)
  الصرف:   المرجع | التاريخ | الجهة (asphalt/concrete/sale/project) | الوحدة | العميل | الكود | الكمية | التكلفة | سعر البيع
            | نوع الجهة | اسم الجهة | غرض الديزل | أطنان الأسفلت
  أطنان الجهات: المرجع | نوع الجهة | اسم الجهة | الأطنان

الاستيراد بيستبدل اللي اتستورد قبل كده للموقع (مواد خام بس): الأرصدة الافتتاحية للمواد الخام، والإنتاج والوارد
والصرف المستورد. اللي اتسجل من جوه البرنامج ما بيتمسحش.
"""
from collections import OrderedDict
from datetime import datetime

from openpyxl import load_workbook

from app.db import SessionLocal
from app.models import (AsphaltPartyTons, CrusherProduction, CrusherProductionLine, Item, OpeningBalance, Receipt,
                        ReceiptLine, Request, RequestLine, Site, Unit)
from scripts.import_excel import clean, num, to_date

MARK = "استيراد مواد خام"


def _rows(wb, name):
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    head = [clean(h) for h in rows[0]]
    return [dict(zip(head, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]


def run(fileobj, site_code):
    wb = load_workbook(fileobj, data_only=True)
    db = SessionLocal()
    try:
        site = db.query(Site).filter_by(code=site_code).first()
        if not site:
            raise ValueError(f"الموقع {site_code} غير موجود")
        # 1) الأصناف
        items = {i.code: i for i in db.query(Item).all()}
        n_new = 0
        for r in _rows(wb, "الأصناف"):
            code = clean(r.get("الكود"))
            if not code:
                continue
            it = items.get(code)
            if not it:
                it = Item(code=code, name=clean(r.get("الاسم")) or code, uom=clean(r.get("الوحدة")) or "طن",
                          category="raw", crusher=bool(num(r.get("ناتج كسارة"))))
                db.add(it)
                db.flush()
                items[code] = it
                n_new += 1

        def item(code):
            it = items.get(clean(code))
            if not it:
                raise ValueError(f"الصنف {code} مش موجود في شيت الأصناف")
            return it

        # 2) مسح الاستيراد القديم للموقع (مواد خام بس)
        raw_ids = [i.id for i in items.values() if i.category == "raw"]
        db.query(OpeningBalance).filter(OpeningBalance.site_id == site.id, OpeningBalance.item_id.in_(raw_ids)) \
            .delete(synchronize_session=False)
        for p in db.query(CrusherProduction).filter_by(site_id=site.id, notes=MARK):
            db.delete(p)
        # فواتير المواد الخام المستوردة (بدون مستخدم وكل سطورها مواد خام) — والمرفقات اللي اترفعت عليها بتترجع بعد الاستيراد
        kept_att = {}
        raw_set = set(raw_ids)
        for rc in db.query(Receipt).filter(Receipt.site_id == site.id, Receipt.created_by_id.is_(None)):
            if rc.lines and all(l.item_id in raw_set for l in rc.lines):
                if rc.attachment and rc.attachment != MARK and rc.invoice_no:
                    kept_att[rc.invoice_no] = rc.attachment
                db.delete(rc)
        for q in db.query(Request).filter_by(site_id=site.id, type="raw", imported=True):
            db.delete(q)
        db.flush()

        # 3) الافتتاحي
        n_ob = 0
        for r in _rows(wb, "الافتتاحي"):
            db.add(OpeningBalance(site_id=site.id, item_id=item(r["الكود"]).id, qty=num(r.get("الكمية")),
                                  value=num(r.get("القيمة")), as_of=to_date(r["التاريخ"]) if r.get("التاريخ") else None))
            n_ob += 1

        # 4) الإنتاج — يوم واحد = تسجيل واحد
        crusher = next((u for u in site.units if u.kind == "crusher" and u.active), None)
        prod = OrderedDict()
        for r in _rows(wb, "الإنتاج"):
            prod.setdefault(to_date(r["التاريخ"]), []).append(r)
        for d, rs in prod.items():
            p = CrusherProduction(site_id=site.id, unit_id=crusher.id if crusher else None, date=d, notes=MARK)
            for r in rs:
                p.lines.append(CrusherProductionLine(item_id=item(r["الكود"]).id, qty=num(r.get("الكمية")),
                                                     unit_price=num(r.get("السعر"))))
            db.add(p)

        # 5) الوارد
        rec = OrderedDict()
        for r in _rows(wb, "الوارد"):
            rec.setdefault((to_date(r["التاريخ"]), clean(r.get("رقم الفاتورة")), clean(r.get("المورد"))), []).append(r)
        for (d, inv, sup), rs in rec.items():
            rc = Receipt(site_id=site.id, date=d, invoice_no=inv, supplier=sup, attachment=kept_att.get(inv),
                         no_vat=bool(num(rs[0].get("بدون ضريبة"))))
            for r in rs:
                rc.lines.append(ReceiptLine(item_id=item(r["الكود"]).id, qty=num(r.get("الكمية")),
                                            unit_price=num(r.get("السعر"))))
            db.add(rc)

        # 6) الصرف — كل مرجع = طلب معتمد
        units = {u.name: u for u in site.units}
        tons = {}
        for r in _rows(wb, "أطنان الجهات"):
            tons.setdefault(clean(r["المرجع"]), []).append(r)
        reqs = OrderedDict()
        for r in _rows(wb, "الصرف"):
            reqs.setdefault(clean(r["المرجع"]), []).append(r)
        for ref, rs in reqs.items():
            f = rs[0]
            dest = clean(f.get("الجهة")) or "asphalt"
            u = units.get(clean(f.get("الوحدة")) or "")
            q = Request(req_no=f"RM-{site.code}-{ref}"[:30], type="raw", site_id=site.id, unit_id=u.id if u else None,
                        requester_name="استيراد", work_date=to_date(f["التاريخ"]), raw_dest=dest,
                        customer=clean(f.get("العميل")), asphalt_tons=num(f.get("أطنان الأسفلت")) or None,
                        status="approved", current_stage=3, imported=True, closed_at=datetime.now(),
                        reason="مستورد من تقارير الإنتاج")
            for r in rs:
                qty, cost = num(r.get("الكمية")), num(r.get("التكلفة"))
                q.lines.append(RequestLine(item_id=item(r["الكود"]).id, qty=qty, status="approved", unit_cost=cost,
                                           value=round(qty * cost, 2), sale_price=num(r.get("سعر البيع")) or None,
                                           entity_type=clean(r.get("نوع الجهة")), entity_name=clean(r.get("اسم الجهة")),
                                           diesel_purpose=clean(r.get("غرض الديزل"))))
            for t in tons.get(ref, []):
                q.party_tons.append(AsphaltPartyTons(entity_type=clean(t.get("نوع الجهة")),
                                                     entity_name=clean(t.get("اسم الجهة")) or "—", tons=num(t.get("الأطنان"))))
            db.add(q)
        db.flush()
        from app.services import revalue_raw
        revalue_raw(db, site.id)   # البيتومين والديزل بالمتوسط المتحرك حسب التاريخ
        db.commit()
        return (f"تم استيراد المواد الخام لموقع {site.name}: أصناف جديدة {n_new}، أرصدة افتتاحية {n_ob}، "
                f"أيام إنتاج {len(prod)}، فواتير وارد {len(rec)}، طلبات صرف {len(reqs)}")
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()
