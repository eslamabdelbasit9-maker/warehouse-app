"""استيراد بيانات موقع من ملف Excel الحالي (منظومة المستودعات) إلى النظام.

الاستخدام:
    python -m scripts.import_excel --file "منظومة_المستودعات_نجران.xlsx" --site-code NJR --site-name "نجران"

يستورد: الأصناف + الرصيد الافتتاحي (شيت الرصيد)، الوارد (شيت الوارد)، الصرف المعتمد والمرفوض (شيت الصرف).
آمن للتشغيل مرة واحدة لكل موقع؛ يرفض التشغيل إذا كان للموقع بيانات سابقة —
إلا لو replace=True: ساعتها يمسح بيانات الاستيراد القديمة للموقع (الأرصدة الافتتاحية، الوارد المستورد،
وطلبات الصرف المستوردة) ويستورد من جديد. اللي اتعمل من جوه البرنامج نفسه (طلبات/وارد/تحويلات/عهد) ما بيتمسحش،
والفواتير المرفوعة جوه البرنامج بترجع تتربط برقم الفاتورة.
"""
import argparse
import re
import sys
from collections import OrderedDict
from datetime import date, datetime

from openpyxl import load_workbook

from app.db import SessionLocal
from app.main import init_db
from app.models import Item, OpeningBalance, Receipt, ReceiptLine, Request, RequestLine, Site, Unit

UNIT_KIND_GUESS = {"كسارة": "crusher", "خلاطة اسفلت": "asphalt", "خلاطة أسفلت": "asphalt", "خرسانة": "concrete"}
_MARKS = re.compile("[\u200e\u200f\u202a-\u202e]")


def clean(s):
    return _MARKS.sub("", str(s)).strip() if s is not None else None


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if v:
        return date.fromisoformat(str(v)[:10])
    return None


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def table_rows(ws, table_name=None):
    if table_name and table_name in ws.tables:
        ref = ws.tables[table_name].ref
        rows = list(ws[ref])
    else:
        rows = list(ws.iter_rows())
    head = [c.value for c in rows[0]]
    for r in rows[1:]:
        vals = [c.value for c in r]
        if any(v not in (None, "") for v in vals):
            yield dict(zip(head, vals))


def inv_key(x):
    return re.sub(r"[^0-9a-z\u0600-\u06ff]", "", str(x or "").lower()).lstrip("0")


def clear_imported(db, site):
    """يمسح بيانات الاستيراد السابقة للموقع، ويرجّع الفواتير المرفوعة جوه البرنامج (رقم الفاتورة ← المرفق)."""
    kept = {}
    old = db.query(Receipt).filter(Receipt.site_id == site.id, Receipt.created_by_id.is_(None)).all()
    for r in old:
        if r.attachment and not r.attachment.startswith("http") and inv_key(r.invoice_no):
            kept.setdefault(inv_key(r.invoice_no), r.attachment)
        db.delete(r)
    n_req = 0
    for q in db.query(Request).filter_by(site_id=site.id, imported=True).all():
        db.delete(q)
        n_req += 1
    n_ob = db.query(OpeningBalance).filter_by(site_id=site.id).delete()
    db.flush()
    return kept, len(old), n_req, n_ob


def run(path, site_code, site_name, replace=False):
    init_db()
    wb = load_workbook(path, data_only=True)
    db = SessionLocal()
    site = db.query(Site).filter_by(code=site_code).first()
    kept = {}
    cleared = None
    if site:
        has_data = db.query(Request).filter_by(site_id=site.id).first() or db.query(OpeningBalance).filter_by(site_id=site.id).first()
        if has_data and not replace:
            db.close()
            raise ValueError(f"الموقع {site_code} فيه بيانات بالفعل — تم الإيقاف لتجنب التكرار (علّم «استبدال» لو عايز تحط نسخة جديدة)")
        if replace:
            kept, *cleared = clear_imported(db, site)
    else:
        site = Site(code=site_code, name=site_name)
        db.add(site)
        db.flush()

    # 1) الأصناف والأرصدة الافتتاحية
    items = {i.code: i for i in db.query(Item).all()}
    n_items = 0
    for r in table_rows(wb["الرصيد"], "BalanceTable"):
        code = clean(r.get("كود الصنف"))
        if not code:
            continue
        it = items.get(code)
        if not it:
            it = Item(code=code, name=clean(r.get("اسم الصنف")) or code, uom=clean(r.get("وحدة القياس")) or "قطعة",
                      category="spare")
            db.add(it)
            db.flush()
            items[code] = it
            n_items += 1
        q, v = num(r.get("الرصيد الافتتاحي")), num(r.get("قيمة الرصيد الافتتاحي"))
        db.add(OpeningBalance(site_id=site.id, item_id=it.id, qty=q, value=v, location=clean(r.get("الموقع"))))

    def item_for(code, name):
        code = clean(code)
        if code not in items:
            it = Item(code=code, name=clean(name) or code, category="spare")
            db.add(it)
            db.flush()
            items[code] = it
        return items[code]

    # 2) الوارد — تجميع حسب (التاريخ، رقم الفاتورة، المورد)
    groups = OrderedDict()
    links = {}
    ws_in = wb["الوارد"]
    head_in = [c.value for c in ws_in[1]]
    link_col = head_in.index("مرفق الفاتورة") if "مرفق الفاتورة" in head_in else None
    for cells in ws_in.iter_rows(min_row=2):
        r = dict(zip(head_in, [c.value for c in cells]))
        if not r.get("كود الصنف"):
            continue
        key = (to_date(r.get("التاريخ")), clean(r.get("رقم الفاتورة")), clean(r.get("المورد")))
        groups.setdefault(key, []).append(r)
        if link_col is not None and key not in links:
            c = cells[link_col]
            url = c.hyperlink.target if c.hyperlink and c.hyperlink.target else c.value
            if isinstance(url, str) and url.startswith("http"):
                links[key] = url[:300]
    for (d, inv, sup), rows in groups.items():
        rec = Receipt(site_id=site.id, date=d or date.today(), invoice_no=inv, supplier=sup,
                      attachment=kept.get(inv_key(inv)) or links.get((d, inv, sup)))
        for r in rows:
            it = item_for(r["كود الصنف"], r.get("اسم الصنف"))
            rec.lines.append(ReceiptLine(item_id=it.id, qty=num(r.get("الكمية")), unit_price=num(r.get("سعر الوحدة"))))
        db.add(rec)

    # 3) الصرف — تجميع حسب ReqID
    units = {u.name: u for u in db.query(Unit).filter_by(site_id=site.id)}
    reqs = OrderedDict()
    for r in table_rows(wb["الصرف"], "IssueTable"):
        if not r.get("كود الصنف"):
            continue
        rid = clean(r.get("ReqID")) or clean(r.get("رقم الطلب"))
        reqs.setdefault(rid, []).append(r)
    n_req = 0
    for rid, rows in reqs.items():
        f = rows[0]
        uname = clean(f.get("المشروع")) or "غير محدد"
        if uname not in units:
            units[uname] = Unit(site_id=site.id, name=uname, kind=UNIT_KIND_GUESS.get(uname, "other"))
            db.add(units[uname])
            db.flush()
        statuses = {clean(r.get("حالة الاعتماد")) for r in rows}
        approved_any = any(s == "معتمد نهائي" for s in statuses)
        if not approved_any and not any((s or "").startswith("مرفوض") for s in statuses):
            # طلب لسه تحت الاعتماد — يدخل النظام كطلب معلّق في المرحلة الأولى
            req = Request(req_no=f"{site_code}-{rid}", type="spare", site_id=site.id, unit_id=units[uname].id,
                          requester_name=clean(f.get("طالب الصرف")), work_date=to_date(f.get("التاريخ")),
                          reason=clean(f.get("السبب")), status="pending", current_stage=1, imported=True)
            for r in rows:
                it = item_for(r["كود الصنف"], r.get("اسم الصنف"))
                req.lines.append(RequestLine(item_id=it.id, qty=num(r.get("الكمية")), status="pending"))
            db.add(req)
            n_req += 1
            continue
        req = Request(req_no=f"{site_code}-{rid}", type="spare", site_id=site.id, unit_id=units[uname].id,
                      requester_name=clean(f.get("طالب الصرف")), work_date=to_date(f.get("التاريخ")),
                      reason=clean(f.get("السبب")), status="approved" if approved_any else "rejected",
                      current_stage=3, imported=True, closed_at=datetime.now())
        for r in rows:
            it = item_for(r["كود الصنف"], r.get("اسم الصنف"))
            ok = clean(r.get("حالة الاعتماد")) == "معتمد نهائي"
            req.lines.append(RequestLine(item_id=it.id, qty=num(r.get("الكمية")), status="approved" if ok else "rejected",
                                         rejected_stage=None if ok else 2,
                                         unit_cost=num(r.get("سعر الوحدة")) if ok else None,
                                         value=num(r.get("القيمة")) if ok else None))
        if approved_any and not all(l.status == "approved" for l in req.lines):
            req.status = "partial"
        db.add(req)
        n_req += 1
    db.commit()
    msg = f"تم: موقع {site.name} — أصناف جديدة {n_items}، فواتير وارد {len(groups)}، طلبات صرف {n_req}، وحدات {len(units)}"
    if cleared:
        msg = (f"اتمسحت النسخة القديمة ({cleared[0]} فاتورة وارد، {cleared[1]} طلب صرف، {cleared[2]} رصيد افتتاحي) — "
               f"واترجع ربط {len(kept)} فاتورة مرفوعة جوه البرنامج. " + msg)
    db.close()
    return msg


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", required=True)
    ap.add_argument("--site-code", required=True)
    ap.add_argument("--site-name", required=True)
    ap.add_argument("--replace", action="store_true")
    a = ap.parse_args()
    try:
        print(run(a.file, a.site_code.upper(), a.site_name, replace=a.replace))
    except ValueError as e:
        sys.exit(str(e))
